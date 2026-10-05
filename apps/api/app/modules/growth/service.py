import csv
import re
from datetime import UTC, datetime, time
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import BadZipFile, ZipFile

import xlrd
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError
from app.infrastructure.db.models import (
    GrowthEvidenceModel,
    MeetingMinutesModel,
    ProgressReportModel,
    TaskModel,
    WorkMaterialModel,
)
from app.modules.growth.schemas import (
    GenerateReportRequest,
    GrowthEvidenceResponse,
    GrowthProfileResponse,
    LearningRecommendation,
    MeetingMinutesResponse,
    ProgressReportResponse,
    SourceReference,
    WorkMaterialResponse,
)

ALLOWED_EXTENSIONS = {
    ".mp3": "audio",
    ".wav": "audio",
    ".m4a": "audio",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".pdf": "document",
    ".docx": "document",
    ".xlsx": "spreadsheet",
    ".xls": "spreadsheet",
    ".csv": "spreadsheet",
    ".txt": "text",
    ".md": "text",
}
EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_PATTERN = re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")
HIGH_RISK_TERMS = ("身份证", "银行卡", "密码", "访问令牌", "access token")
PDF_MAX_PAGES = 50
PDF_MAX_EXTRACTED_CHARS = 50_000
DOCUMENT_MAX_EXTRACTED_CHARS = 50_000
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class GrowthService:
    def __init__(self, session: AsyncSession, max_upload_bytes: int) -> None:
        self.session = session
        self.max_upload_bytes = max_upload_bytes

    async def create_material(
        self,
        user_id: UUID,
        filename: str,
        mime_type: str,
        content: bytes,
        purpose: str,
    ) -> WorkMaterialResponse:
        extension = Path(filename).suffix.lower()
        material_type = ALLOWED_EXTENSIONS.get(extension)
        if material_type is None:
            raise AppError(
                "UNSUPPORTED_MATERIAL_TYPE",
                "仅支持音频、截图、PDF、DOCX、Excel、CSV、Markdown 和 TXT 材料",
                status_code=415,
            )
        if len(content) > self.max_upload_bytes:
            raise AppError("FILE_TOO_LARGE", "文件不能超过 20 MB", status_code=413)

        # TODO(multimodal-privacy): Run OCR/privacy detection on images before they can be
        # attached to a chat and sent to an external vision provider.
        extracted = self._extract_text(extension, content)
        redacted, privacy_status = self._redact(extracted)
        status = "ready" if redacted else "needs_confirmation"
        if privacy_status == "review_required":
            status = "blocked"
        material = WorkMaterialModel(
            id=str(uuid4()),
            user_id=str(user_id),
            title=Path(filename).name,
            material_type=material_type,
            mime_type=mime_type or "application/octet-stream",
            size_bytes=len(content),
            purpose=purpose,
            status=status,
            privacy_status=privacy_status,
            content_excerpt=redacted[:4000],
            extracted_text=redacted[:DOCUMENT_MAX_EXTRACTED_CHARS],
            binary_content=(
                content
                if material_type == "image" or extension in {".pdf", ".docx", ".xlsx"}
                else None
            ),
            created_at=self._now(),
        )
        self.session.add(material)
        await self.session.flush()

        minutes: MeetingMinutesModel | None = None
        if purpose == "meeting" and status != "blocked":
            minutes = self._build_minutes(user_id, material, redacted)
            self.session.add(minutes)
            evidence = self._build_material_evidence(user_id, material, redacted)
            self.session.add(evidence)
        await self.session.commit()
        return self._material_response(material, minutes)

    async def list_materials(self, user_id: UUID) -> list[WorkMaterialResponse]:
        materials = list(
            await self.session.scalars(
                select(WorkMaterialModel)
                .where(WorkMaterialModel.user_id == str(user_id))
                .order_by(WorkMaterialModel.created_at.desc())
            )
        )
        minutes_by_source = {
            item.source_material_id: item
            for item in await self.session.scalars(
                select(MeetingMinutesModel).where(MeetingMinutesModel.user_id == str(user_id))
            )
        }
        return [self._material_response(item, minutes_by_source.get(item.id)) for item in materials]

    async def get_material(self, user_id: UUID, material_id: UUID) -> WorkMaterialResponse:
        material = await self._owned_material(user_id, material_id)
        minutes = await self.session.scalar(
            select(MeetingMinutesModel).where(
                MeetingMinutesModel.source_material_id == material.id,
                MeetingMinutesModel.user_id == str(user_id),
            )
        )
        return self._material_response(material, minutes)

    async def delete_material(self, user_id: UUID, material_id: UUID) -> None:
        material = await self._owned_material(user_id, material_id)
        await self.session.execute(
            delete(MeetingMinutesModel).where(
                MeetingMinutesModel.user_id == str(user_id),
                MeetingMinutesModel.source_material_id == material.id,
            )
        )
        await self.session.execute(
            delete(GrowthEvidenceModel).where(
                GrowthEvidenceModel.user_id == str(user_id),
                GrowthEvidenceModel.source_type == "material",
                GrowthEvidenceModel.source_id == material.id,
            )
        )
        reports = list(
            await self.session.scalars(
                select(ProgressReportModel).where(ProgressReportModel.user_id == str(user_id))
            )
        )
        for report in reports:
            if any(
                source.get("type") == "material" and source.get("id") == material.id
                for source in report.source_refs
            ):
                await self.session.delete(report)
        await self.session.delete(material)
        await self.session.commit()

    async def generate_report(
        self, user_id: UUID, payload: GenerateReportRequest
    ) -> ProgressReportResponse:
        start_dt = datetime.combine(payload.period_start, time.min)
        end_dt = datetime.combine(payload.period_end, time.max)
        tasks = list(
            await self.session.scalars(
                select(TaskModel)
                .where(
                    TaskModel.user_id == str(user_id),
                    TaskModel.status == "done",
                    TaskModel.completed_at >= start_dt,
                    TaskModel.completed_at <= end_dt,
                )
                .order_by(TaskModel.completed_at.asc())
            )
        )
        materials = list(
            await self.session.scalars(
                select(WorkMaterialModel)
                .where(
                    WorkMaterialModel.user_id == str(user_id),
                    WorkMaterialModel.status == "ready",
                    WorkMaterialModel.created_at >= start_dt,
                    WorkMaterialModel.created_at <= end_dt,
                )
                .order_by(WorkMaterialModel.created_at.asc())
            )
        )

        completed = [f"完成「{task.title}」：{task.done_definition}" for task in tasks]
        learned = [
            f"从「{item.title}」沉淀：{item.content_excerpt[:120]}"
            for item in materials
            if item.content_excerpt
        ]
        sections = {
            "completed": completed or ["本周期暂无可验证的已完成事项。"],
            "learnings": learned or ["本周期暂无可追溯的材料沉淀。"],
            "risks": ["请确认草稿中的事实与敏感信息后再对外发送。"],
            "next_steps": ["结合行动计划补充下一周期目标，并为每项结果保留来源。"],
        }
        source_refs = [{"type": "task", "id": task.id, "title": task.title} for task in tasks] + [
            {"type": "material", "id": item.id, "title": item.title} for item in materials
        ]
        label = "周报" if payload.period_type == "weekly" else "月报"
        report = ProgressReportModel(
            id=str(uuid4()),
            user_id=str(user_id),
            period_type=payload.period_type,
            period_start=payload.period_start,
            period_end=payload.period_end,
            title=f"{payload.period_start:%m.%d}-{payload.period_end:%m.%d} {label}草稿",
            sections=sections,
            source_refs=source_refs,
            created_at=self._now(),
        )
        self.session.add(report)
        for task in tasks:
            exists = await self.session.scalar(
                select(GrowthEvidenceModel.id).where(
                    GrowthEvidenceModel.user_id == str(user_id),
                    GrowthEvidenceModel.source_type == "task",
                    GrowthEvidenceModel.source_id == task.id,
                )
            )
            if exists is None:
                self.session.add(
                    GrowthEvidenceModel(
                        id=str(uuid4()),
                        user_id=str(user_id),
                        capability=self._capability_for_text(task.title + task.description),
                        summary=f"按完成定义交付：{task.done_definition}",
                        source_type="task",
                        source_id=task.id,
                        source_title=task.title,
                        observed_at=task.completed_at or self._now(),
                    )
                )
        await self.session.commit()
        return self._report_response(report)

    async def list_reports(self, user_id: UUID) -> list[ProgressReportResponse]:
        reports = await self.session.scalars(
            select(ProgressReportModel)
            .where(ProgressReportModel.user_id == str(user_id))
            .order_by(ProgressReportModel.created_at.desc())
        )
        return [self._report_response(item) for item in reports]

    async def growth_profile(self, user_id: UUID) -> GrowthProfileResponse:
        evidence = list(
            await self.session.scalars(
                select(GrowthEvidenceModel)
                .where(GrowthEvidenceModel.user_id == str(user_id))
                .order_by(GrowthEvidenceModel.observed_at.desc())
            )
        )
        counts: dict[str, int] = {}
        for item in evidence:
            counts[item.capability] = counts.get(item.capability, 0) + 1
        focus = sorted(counts.items(), key=lambda item: (-item[1], item[0]))
        recommendations = [
            LearningRecommendation(
                capability=capability,
                reason=f"已有 {count} 条可追溯证据，可通过刻意练习继续巩固。",
                next_action=f"下一次相关任务后记录结果与反馈，补充一条「{capability}」证据。",
                evidence_count=count,
            )
            for capability, count in focus[:3]
        ]
        if not recommendations:
            recommendations = [
                LearningRecommendation(
                    capability="工作复盘",
                    reason="目前还没有可追溯证据。",
                    next_action="完成一项计划任务或导入一份会议材料，建立第一条成长证据。",
                    evidence_count=0,
                )
            ]
        return GrowthProfileResponse(
            evidence=[self._evidence_response(item) for item in evidence],
            recommendations=recommendations,
        )

    async def _owned_material(self, user_id: UUID, material_id: UUID) -> WorkMaterialModel:
        material = await self.session.scalar(
            select(WorkMaterialModel).where(
                WorkMaterialModel.id == str(material_id),
                WorkMaterialModel.user_id == str(user_id),
            )
        )
        if material is None:
            raise AppError("MATERIAL_NOT_FOUND", "资料不存在", status_code=404)
        return material

    @staticmethod
    def _extract_text(extension: str, content: bytes) -> str:
        if extension in {".txt", ".md"}:
            return content.decode("utf-8", errors="replace").strip()
        if extension == ".csv":
            return GrowthService._extract_csv(content)
        if extension == ".xlsx":
            return GrowthService._extract_xlsx(content)
        if extension == ".xls":
            return GrowthService._extract_xls(content)
        if extension == ".pdf":
            return GrowthService._extract_pdf(content)
        if extension == ".docx":
            return GrowthService._extract_docx(content)
        if extension in {".mp3", ".wav", ".m4a"}:
            # TODO(multimodal-audio): Add speech-to-text transcription before audio
            # attachments are advertised as readable chat context.
            return ""
        return ""

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        """Extract visible text from OOXML parts without treating XML markup as text."""
        parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
        try:
            with ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
                if "word/document.xml" not in names:
                    raise KeyError("word/document.xml")
                part_names = ["word/document.xml"]
                part_names.extend(
                    sorted(
                        name
                        for name in names
                        if name.startswith(("word/header", "word/footer"))
                        and name.endswith(".xml")
                    )
                )
                part_names.extend(
                    name
                    for name in ("word/footnotes.xml", "word/endnotes.xml")
                    if name in names
                )

                sections: list[str] = []
                paragraph_tag = f"{{{WORD_NAMESPACE}}}p"
                text_tag = f"{{{WORD_NAMESPACE}}}t"
                for part_name in part_names:
                    root = etree.fromstring(archive.read(part_name), parser=parser)
                    for paragraph in root.iter(paragraph_tag):
                        # Text boxes contain nested paragraphs. Only leaf paragraphs are emitted
                        # so their text is not duplicated by an outer drawing paragraph.
                        if any(
                            descendant is not paragraph
                            for descendant in paragraph.iter(paragraph_tag)
                        ):
                            continue
                        text = "".join(
                            node.text or "" for node in paragraph.iter(text_tag)
                        ).strip()
                        if text:
                            sections.append(text)
                return "\n".join(sections).strip()
        except (BadZipFile, KeyError, etree.XMLSyntaxError, ValueError) as exc:
            raise AppError(
                "DOCX_PARSE_FAILED", "Word 文件无法读取或已损坏", status_code=422
            ) from exc

    @staticmethod
    def _extract_pdf(content: bytes) -> str:
        try:
            reader = PdfReader(BytesIO(content), strict=False)
            if reader.is_encrypted and not reader.decrypt(""):
                raise AppError(
                    "PDF_PASSWORD_REQUIRED",
                    "PDF 文件已加密，请上传未加密版本",
                    status_code=422,
                )
            sections: list[str] = []
            extracted_chars = 0
            for page_number, page in enumerate(reader.pages, start=1):
                if page_number > PDF_MAX_PAGES or extracted_chars >= PDF_MAX_EXTRACTED_CHARS:
                    break
                text = (page.extract_text() or "").strip()
                if not text:
                    continue
                remaining = PDF_MAX_EXTRACTED_CHARS - extracted_chars
                section = f"[PDF 第 {page_number} 页]\n{text[:remaining]}"
                sections.append(section)
                extracted_chars += len(section)
            return "\n\n".join(sections)
        except AppError:
            raise
        except (PdfReadError, OSError, TypeError, ValueError) as exc:
            raise AppError("PDF_PARSE_FAILED", "PDF 文件无法读取或已损坏", status_code=422) from exc

    @staticmethod
    def _extract_csv(content: bytes) -> str:
        decoded: str | None = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise AppError("SPREADSHEET_PARSE_FAILED", "无法识别 CSV 文件编码", status_code=422)
        return GrowthService._format_sheet("CSV", list(csv.reader(decoded.splitlines()))[:50])

    @staticmethod
    def _extract_xlsx(content: bytes) -> str:
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
            raise AppError(
                "SPREADSHEET_PARSE_FAILED", "Excel 文件无法读取或已损坏", status_code=422
            ) from exc
        sections: list[str] = []
        try:
            for sheet in workbook.worksheets[:5]:
                rows: list[list[object]] = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 50:
                        break
                    rows.append(list(row[:20]))
                formatted = GrowthService._format_sheet(sheet.title, rows)
                if formatted:
                    sections.append(formatted)
        finally:
            workbook.close()
        return "\n".join(sections)

    @staticmethod
    def _extract_xls(content: bytes) -> str:
        try:
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        except (xlrd.biffh.XLRDError, ValueError) as exc:
            raise AppError(
                "SPREADSHEET_PARSE_FAILED", "Excel 文件无法读取或已损坏", status_code=422
            ) from exc
        sections: list[str] = []
        try:
            for sheet_name in workbook.sheet_names()[:5]:
                sheet = workbook.sheet_by_name(sheet_name)
                rows = [
                    [
                        sheet.cell_value(row_index, column_index)
                        for column_index in range(min(sheet.ncols, 20))
                    ]
                    for row_index in range(min(sheet.nrows, 50))
                ]
                formatted = GrowthService._format_sheet(sheet_name, rows)
                if formatted:
                    sections.append(formatted)
        finally:
            workbook.release_resources()
        return "\n".join(sections)

    @staticmethod
    def _format_sheet(sheet_name: str, rows: list[list[object]]) -> str:
        lines = [f"[工作表：{sheet_name}]"]
        has_values = False
        for index, row in enumerate(rows):
            values = [
                str(value).replace("\r", " ").replace("\n", " ").strip()[:120]
                if value is not None
                else ""
                for value in row[:20]
            ]
            while values and not values[-1]:
                values.pop()
            if not any(values):
                continue
            has_values = True
            label = "表头" if index == 0 else f"第 {index + 1} 行"
            lines.append(f"{label}：" + " | ".join(values))
        return "\n".join(lines) if has_values else ""

    @staticmethod
    def _redact(text: str) -> tuple[str, str]:
        if any(term.lower() in text.lower() for term in HIGH_RISK_TERMS):
            return "材料包含高风险隐私字段，已阻止自动处理。", "review_required"
        redacted = EMAIL_PATTERN.sub("[邮箱已脱敏]", text)
        redacted = PHONE_PATTERN.sub("[手机号已脱敏]", redacted)
        return redacted, "redacted" if redacted != text else "clear"

    def _build_minutes(
        self, user_id: UUID, material: WorkMaterialModel, text: str
    ) -> MeetingMinutesModel:
        if not text:
            summary = (
                "材料已接收；当前未提取到可读取文字，且未配置音频/图像识别，"
                "请补充文字纪要后再确认事实。"
            )
            return MeetingMinutesModel(
                id=str(uuid4()),
                user_id=str(user_id),
                source_material_id=material.id,
                summary=summary,
                action_items=[],
                pending_facts=["待补充可验证的文字内容"],
                created_at=self._now(),
            )
        lines = [line.strip(" -#\t") for line in text.splitlines() if line.strip()]
        action_items = [
            line
            for line in lines
            if any(key in line.lower() for key in ("待办", "行动", "todo", "负责", "完成"))
        ][:8]
        pending = [
            line
            for line in lines
            if any(key in line.lower() for key in ("待确认", "不确定", "?", "？"))
        ][:8]
        summary_lines = [line for line in lines if line not in action_items and line not in pending]
        summary = "；".join(summary_lines[:3])[:1000] or "已提取会议材料，请人工确认以下事实。"
        return MeetingMinutesModel(
            id=str(uuid4()),
            user_id=str(user_id),
            source_material_id=material.id,
            summary=summary,
            action_items=action_items,
            pending_facts=pending or ["请确认纪要摘要与行动项是否完整"],
            created_at=self._now(),
        )

    def _build_material_evidence(
        self, user_id: UUID, material: WorkMaterialModel, text: str
    ) -> GrowthEvidenceModel:
        return GrowthEvidenceModel(
            id=str(uuid4()),
            user_id=str(user_id),
            capability=self._capability_for_text(text),
            summary=(text[:500] or "完成材料整理，内容等待人工确认。"),
            source_type="material",
            source_id=material.id,
            source_title=material.title,
            observed_at=material.created_at,
        )

    @staticmethod
    def _capability_for_text(text: str) -> str:
        if any(key in text for key in ("汇报", "周报", "表达", "沟通")):
            return "沟通与汇报"
        if any(key.lower() in text.lower() for key in ("react", "代码", "开发", "技术")):
            return "专业技能"
        if any(key in text for key in ("需求", "项目", "计划", "交付")):
            return "项目执行"
        return "协作与复盘"

    @staticmethod
    def _now() -> datetime:
        return datetime.now(UTC).replace(tzinfo=None)

    @staticmethod
    def _minutes_response(item: MeetingMinutesModel) -> MeetingMinutesResponse:
        return MeetingMinutesResponse(
            id=item.id,
            source_material_id=item.source_material_id,
            summary=item.summary,
            action_items=item.action_items,
            pending_facts=item.pending_facts,
            created_at=item.created_at,
        )

    @classmethod
    def _material_response(
        cls, item: WorkMaterialModel, minutes: MeetingMinutesModel | None
    ) -> WorkMaterialResponse:
        return WorkMaterialResponse(
            id=item.id,
            title=item.title,
            material_type=item.material_type,
            mime_type=item.mime_type,
            size_bytes=item.size_bytes,
            purpose=item.purpose,
            status=item.status,
            privacy_status=item.privacy_status,
            content_excerpt=item.content_excerpt,
            created_at=item.created_at,
            minutes=cls._minutes_response(minutes) if minutes else None,
        )

    @staticmethod
    def _report_response(item: ProgressReportModel) -> ProgressReportResponse:
        return ProgressReportResponse(
            id=item.id,
            period_type=item.period_type,
            period_start=item.period_start,
            period_end=item.period_end,
            title=item.title,
            sections=item.sections,
            sources=[SourceReference.model_validate(source) for source in item.source_refs],
            created_at=item.created_at,
        )

    @staticmethod
    def _evidence_response(item: GrowthEvidenceModel) -> GrowthEvidenceResponse:
        return GrowthEvidenceResponse(
            id=item.id,
            capability=item.capability,
            summary=item.summary,
            source_type=item.source_type,
            source_id=item.source_id,
            source_title=item.source_title,
            observed_at=item.observed_at,
        )
