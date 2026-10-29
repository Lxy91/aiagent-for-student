import re
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError
from app.infrastructure.db.models import GeneratedArtifactModel

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DOCX_MARKERS = (
    "生成文档",
    "导出文档",
    "下载文档",
    "生成word",
    "导出word",
    "下载word",
    "word文档",
    "docx",
)
XLSX_MARKERS = (
    "生成excel",
    "导出excel",
    "下载excel",
    "excel表格",
    "excel文件",
    "导出表格",
    "下载表格",
    "xlsx",
)
DOCX_CREATION_PATTERN = re.compile(
    r"(?:写|撰写|整理|制作|创建|生成|输出)(?:成|为)?(?:一份|一个|份|个)(?:word|docx)"
    r"(?:报告|总结|方案|简历|材料|文件|文档)?"
)


def requested_artifact_types(content: str) -> tuple[str, ...]:
    normalized = re.sub(r"\s+", "", content).lower()
    result: list[str] = []
    if any(marker in normalized for marker in DOCX_MARKERS) or DOCX_CREATION_PATTERN.search(
        normalized
    ):
        result.append("docx")
    if any(marker in normalized for marker in XLSX_MARKERS):
        result.append("xlsx")
    return tuple(result)


def artifact_generation_instruction(content: str) -> str:
    artifact_types = requested_artifact_types(content)
    if not artifact_types:
        return ""
    labels = ["Word（DOCX）" if item == "docx" else "Excel（XLSX）" for item in artifact_types]
    file_label = "、".join(labels)
    return (
        f"当前系统具备生成并提供可下载的 {file_label}文件的能力。"
        "用户已明确要求创作文件，请直接撰写适合写入文件的完整内容；"
        "正文第一行必须用一级 Markdown 标题‘# 文档标题’给出简洁、正式的"
        "文件名称。标题应根据产物类型命名，例如‘个人简历分析报告’、"
        "‘前端岗位面试自我介绍稿’或‘项目进度跟踪表’，不得复制或改写用户问题。"
        "系统会使用该一级标题作为下载文件名。"
        "回复完成后，系统会自动生成文件并展示下载链接。"
        "不得声称无法生成、保存或下载文件，也不要让用户复制到本地软件自行创建。"
        "如果用户尚未指定主题或具体内容，请生成一份简洁、可继续编辑的通用模板，"
        "并说明用户可以继续补充内容进行修改。"
    )


def _sanitize_title(title: str) -> str:
    title = re.sub(r"[*_`~]", "", title)
    title = re.sub(r"[\\/:*?\"<>|\r\n]+", " ", title)
    title = " ".join(title.split()).strip(" ，。；：-")
    return title[:48]


def _document_title(content: str, artifact_type: str) -> str:
    """Use the AI-authored H1 as the filename, with a semantic fallback."""
    for line in content.splitlines():
        heading = re.match(r"^\s*#\s+(.+?)\s*$", line)
        if heading:
            title = _sanitize_title(heading.group(1))
            if title:
                return title

    semantic_titles = (
        (("简历",), "个人简历分析报告"),
        (("自我介绍", "面试"), "面试自我介绍稿"),
        (("周报",), "工作周报"),
        (("月报",), "工作月报"),
        (("职业规划",), "职业发展规划报告"),
        (("复盘",), "工作复盘报告"),
        (("调研",), "调研分析报告"),
        (("方案",), "工作方案"),
    )
    for markers, title in semantic_titles:
        if all(marker in content for marker in markers):
            return title
    return "数据整理表" if artifact_type == "xlsx" else "智能对话整理报告"


def _document_lines(content: str, title: str) -> list[str]:
    lines = content.splitlines()
    for index, line in enumerate(lines):
        heading = re.match(r"^\s*#\s+(.+?)\s*$", line)
        if heading and _sanitize_title(heading.group(1)) == title:
            return lines[:index] + lines[index + 1 :]
        if line.strip():
            break
    return lines


def _parse_markdown_table(lines: list[str]) -> tuple[list[str], list[list[str]]] | None:
    for index in range(len(lines) - 1):
        first = lines[index].strip()
        separator = lines[index + 1].strip()
        if "|" not in first or not re.match(r"^\s*\|?\s*:?-{3,}", separator):
            continue
        headers = [cell.strip() for cell in first.strip("|").split("|")]
        rows: list[list[str]] = []
        for line in lines[index + 2 :]:
            if "|" not in line:
                break
            row = [cell.strip() for cell in line.strip().strip("|").split("|")]
            rows.append((row + [""] * len(headers))[: len(headers)])
        if headers:
            return headers, rows
    return None


def build_docx(prompt: str, content: str) -> tuple[str, bytes]:
    from app.modules.artifacts.docx_builder import render_docx

    del prompt
    title = _document_title(content, "docx")
    body = "\n".join(_document_lines(content, title))
    return f"{title}.docx", render_docx(title, body)


def build_xlsx(prompt: str, content: str) -> tuple[str, bytes]:
    del prompt
    title = _document_title(content, "xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "内容"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A3"

    sheet.merge_cells("A1:C1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="等线", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="237A68")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    lines = _document_lines(content, title)
    table_data = _parse_markdown_table(lines)
    if table_data:
        headers, rows = table_data
        matrix = [headers, *rows]
    else:
        content_rows = [
            [index, line.strip()] for index, line in enumerate(lines, 1) if line.strip()
        ]
        matrix = [["序号", "内容"], *content_rows]

    start_row = 3
    for row_index, row in enumerate(matrix, start_row):
        for column_index, value in enumerate(row, 1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = Font(name="等线", size=10, bold=row_index == start_row)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index == start_row:
                cell.fill = PatternFill("solid", fgColor="DDEDE8")
                cell.font = Font(name="等线", size=10, bold=True, color="185B4E")
    last_row = start_row + len(matrix) - 1
    last_col = max((len(row) for row in matrix), default=1)
    light = Side(style="thin", color="D9E3DF")
    for row in sheet.iter_rows(min_row=start_row, max_row=last_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.border = Border(bottom=light)
    for column_index in range(1, last_col + 1):
        values = [
            str(sheet.cell(row=row, column=column_index).value or "")
            for row in range(start_row, last_row + 1)
        ]
        longest = max((len(value) for value in values), default=10)
        column_letter = get_column_letter(column_index)
        sheet.column_dimensions[column_letter].width = min(max(longest + 3, 12), 48)
    sheet.auto_filter.ref = f"A{start_row}:{sheet.cell(row=last_row, column=last_col).coordinate}"

    payload = BytesIO()
    workbook.save(payload)
    return f"{title}.xlsx", payload.getvalue()


class ArtifactService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def create_requested(
        self,
        *,
        user_id: UUID,
        conversation_id: UUID,
        prompt: str,
        content: str,
    ) -> list[dict]:
        artifacts: list[dict] = []
        for artifact_type in requested_artifact_types(prompt):
            if artifact_type == "docx":
                filename, binary = build_docx(prompt, content)
                mime_type = DOCX_MIME
            else:
                filename, binary = build_xlsx(prompt, content)
                mime_type = XLSX_MIME
            model = GeneratedArtifactModel(
                id=str(uuid4()),
                user_id=str(user_id),
                conversation_id=str(conversation_id),
                filename=filename,
                artifact_type=artifact_type,
                mime_type=mime_type,
                size_bytes=len(binary),
                binary_content=binary,
            )
            self.session.add(model)
            artifacts.append(
                {
                    "id": model.id,
                    "filename": model.filename,
                    "artifact_type": model.artifact_type,
                    "mime_type": model.mime_type,
                    "size_bytes": model.size_bytes,
                }
            )
        if artifacts:
            await self.session.commit()
        return artifacts

    async def create_binary(
        self,
        *,
        user_id: UUID,
        conversation_id: UUID,
        filename: str,
        artifact_type: str,
        mime_type: str,
        binary: bytes,
    ) -> dict:
        model = GeneratedArtifactModel(
            id=str(uuid4()),
            user_id=str(user_id),
            conversation_id=str(conversation_id),
            filename=filename,
            artifact_type=artifact_type,
            mime_type=mime_type,
            size_bytes=len(binary),
            binary_content=binary,
        )
        self.session.add(model)
        await self.session.commit()
        return {
            "id": model.id,
            "filename": model.filename,
            "artifact_type": model.artifact_type,
            "mime_type": model.mime_type,
            "size_bytes": model.size_bytes,
        }

    async def get_owned(self, user_id: UUID, artifact_id: UUID) -> GeneratedArtifactModel:
        artifact = await self.session.scalar(
            select(GeneratedArtifactModel).where(
                GeneratedArtifactModel.id == str(artifact_id),
                GeneratedArtifactModel.user_id == str(user_id),
            )
        )
        if artifact is None:
            raise AppError("ARTIFACT_NOT_FOUND", "文件不存在或无权访问", status_code=404)
        return artifact
