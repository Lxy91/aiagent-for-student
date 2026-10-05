import json
import re
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError
from app.infrastructure.db.models import GeneratedArtifactModel, WorkMaterialModel
from app.modules.artifacts.service import DOCX_MIME, XLSX_MIME, ArtifactService

EDIT_INTENT_MARKERS = (
    "补全",
    "填写",
    "填一下",
    "修改",
    "改一下",
    "完善",
    "更新",
    "编辑",
    "润色",
    "按模板",
)
EDITABLE_EXTENSIONS = {".docx", ".xlsx"}
FORMAT_INTENT_MARKERS = ("统一", "格式", "排版", "字体", "颜色")
ALLOW_BLANK_PATTERNS = (
    re.compile(
        r"(?:不清楚|不知道|不确定|没有|缺少|无法提供|未提供).{0,16}"
        r"(?:不填|不填写|不用填|无需填|留空|空着)"
    ),
    re.compile(r"(?:不填|不填写|不用填|无需填|留空|空着).{0,10}(?:即可|就行|可以)"),
)
EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_PATTERN = re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")
DOCX_LOCATION = re.compile(r"^table:(\d+):row:(\d+):cell:(\d+)$")
DOCX_PARAGRAPH_LOCATION = re.compile(r"^paragraph:(\d+)$")
EXCEL_CELL = re.compile(r"^[A-Z]{1,3}[1-9]\d{0,6}$")
EDIT_ACTION_PATTERN = r"(?:补全|补充|填写|修改|完善|更新|编辑)"
EDIT_TARGET_AFTER_ACTION = re.compile(
    rf"{EDIT_ACTION_PATTERN}.{{0,8}}(?:文档|文件|附件)([一二两12])"
)
EDIT_TARGET_BEFORE_ACTION = re.compile(
    rf"(?:文档|文件|附件)([一二两12]).{{0,8}}{EDIT_ACTION_PATTERN}"
)


def editable_attachments(content: str, attachments: list[dict[str, str]]) -> list[dict[str, str]]:
    normalized = re.sub(r"\s+", "", content).lower()
    if not any(marker in normalized for marker in (*EDIT_INTENT_MARKERS, *FORMAT_INTENT_MARKERS)):
        return []
    return [
        item
        for item in attachments
        if Path(item["title"]).suffix.lower() in EDITABLE_EXTENSIONS
    ]


def select_edit_target(
    content: str, candidates: list[dict[str, str]]
) -> dict[str, str] | None:
    """Select an explicitly numbered target while keeping other files as references."""
    if len(candidates) == 1:
        return candidates[0]
    normalized = re.sub(r"\s+", "", content).lower()
    match = EDIT_TARGET_AFTER_ACTION.search(normalized)
    if not match:
        match = EDIT_TARGET_BEFORE_ACTION.search(normalized)
    if not match:
        return None
    token = match.group(1)
    index = 0 if token in {"一", "1"} else 1
    return candidates[index] if index < len(candidates) else None


def allows_unfilled_fields(content: str) -> bool:
    normalized = re.sub(r"\s+", "", content).lower()
    return any(pattern.search(normalized) for pattern in ALLOW_BLANK_PATTERNS)


def requested_docx_formatting(content: str, filename: str) -> dict[str, str]:
    normalized = re.sub(r"\s+", "", content).lower()
    if Path(filename).suffix.lower() != ".docx":
        return {}
    if "字体" in normalized and "颜色" in normalized and "统一" in normalized:
        return {"font_color": "000000"}
    return {}


def build_edit_plan_messages(
    *, user_request: str, filename: str, descriptor: str, conversation_context: str = ""
) -> list[dict[str, str]]:
    today = date.today().isoformat()
    return [
        {
            "role": "system",
            "content": (
                "你是文件编辑规划器，只输出 JSON 对象。根据用户要求和文件结构生成最小修改计划。"
                "严禁编造单位名称、日期、签名、印章、联系方式、金额等客观事实；缺少这些信息时放入"
                " missing_information。主观问卷、总结、感受类字段可以依据附件和对话内容起草，但不得"
                "冒充单位评价。changes 中 DOCX 使用 location、expected_text、value；XLSX 使用 "
                "sheet、cell、expected_text、value。location/cell 必须来自文件结构，expected_text "
                "必须与当前"
                "内容一致。不要修改公式，不要提出文件中不存在的位置。返回结构："
                "如果用户明确表示未知、不清楚或无法提供的字段可以留空，就不要把这些字段放入"
                " missing_information，应保留原样并继续生成其余安全修改。"
                "文件结构中的非空字段视为已填写，不要再次索要，也不要生成与当前值相同的 change。"
                "<空> 只是空白位置标记，绝不能作为 value 写入文件；保持空白时不要生成 change。"
                "expected_text 必须逐字复制文件结构中该位置的完整当前值；文件结构用 JSON 字符串"
                "表示换行和空格。对于同时包含标题、填写区、签名或日期的单元格，value 必须保留"
                "原有标题、签名和日期文字，只在填写区加入内容。包含[邮箱已脱敏]或[手机号已脱敏]"
                "的现有字段不得修改。"
                f"系统当前日期是 {today}；用户说‘今天’时使用这个日期，不要要求再次确认。"
                '{"message":"给用户的简短说明","missing_information":[],"changes":[]}'
            ),
        },
        {
            "role": "user",
            "content": (
                f"文件名：{filename}\n用户要求：{user_request}\n\n"
                f"近期对话：\n{conversation_context[-6000:]}\n\n"
                f"文件结构（<空> 表示可填写位置）：\n{descriptor}"
            ),
        },
    ]


def normalize_edit_plan(payload: dict[str, Any]) -> tuple[str, list[str], list[dict[str, str]]]:
    message = str(payload.get("message") or "").strip()
    raw_missing = payload.get("missing_information", [])
    if not isinstance(raw_missing, list):
        raw_missing = []
    missing = [
        str(item).strip()
        for item in raw_missing
        if str(item).strip()
    ][:20]
    changes: list[dict[str, str]] = []
    raw_changes = payload.get("changes", [])
    if not isinstance(raw_changes, list):
        raw_changes = []
    for raw in raw_changes[:100]:
        if not isinstance(raw, dict):
            continue
        value = str(raw.get("value") or "").strip()
        if not value or value == "<空>":
            continue
        change = {
            key: str(raw.get(key) or "").strip()
            for key in ("location", "sheet", "cell", "expected_text", "value")
        }
        if change["expected_text"] == "<空>":
            change["expected_text"] = ""
        changes.append(change)
    return message, missing, changes


def _edited_filename(filename: str, action: str) -> str:
    path = Path(filename)
    if "字体" in action and "颜色" in action:
        suffix = "黑色"
    else:
        suffix = "已补全" if "补全" in action or "填写" in action else "已修改"
    return f"{path.stem}_{suffix}{path.suffix.lower()}"


def _write_paragraph(paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _write_cell(cell, value: str) -> None:
    """Replace cell text while preserving form anchors and response spacing."""
    existing_nonempty = [
        (index, paragraph.text.strip())
        for index, paragraph in enumerate(cell.paragraphs)
        if paragraph.text.strip()
    ]
    target_nonempty = [line.strip() for line in value.splitlines() if line.strip()]
    if (
        len(existing_nonempty) >= 2
        and len(target_nonempty) >= 2
        and _comparable_text(existing_nonempty[0][1]) == _comparable_text(target_nonempty[0])
        and _comparable_text(existing_nonempty[-1][1]) == _comparable_text(target_nonempty[-1])
    ):
        first_index = existing_nonempty[0][0]
        last_index = existing_nonempty[-1][0]
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.text = ""
        _write_paragraph(cell.paragraphs[first_index], target_nonempty[0])
        middle = "\n".join(target_nonempty[1:-1])
        if middle:
            response_index = min(first_index + 1, last_index - 1)
            _write_paragraph(cell.paragraphs[response_index], middle)
        _write_paragraph(cell.paragraphs[last_index], target_nonempty[-1])
        return

    _write_paragraph(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def _comparable_text(value: str) -> str:
    # Models commonly compress template whitespace. The exact location plus the
    # whitespace-insensitive original text still protects against stale writes.
    return re.sub(r"\s+", "", value)


def _describe_docx(binary: bytes) -> str:
    document = Document(BytesIO(binary))
    lines: list[str] = []
    for index, paragraph in enumerate(document.paragraphs[:120]):
        text = paragraph.text.strip()
        if text:
            lines.append(f"paragraph:{index} = {text[:300]}")
    cell_count = 0
    for table_index, table in enumerate(document.tables[:20]):
        seen_cells: set[Any] = set()
        for row_index, row in enumerate(table.rows[:80]):
            for cell_index, cell in enumerate(row.cells[:20]):
                cell_identity = cell._tc
                if cell_identity in seen_cells:
                    continue
                seen_cells.add(cell_identity)
                value = cell.text.strip() or "<空>"
                lines.append(
                    f"table:{table_index}:row:{row_index}:cell:{cell_index} = "
                    f"{json.dumps(value[:300], ensure_ascii=False)}"
                )
                cell_count += 1
                if cell_count >= 500:
                    return "\n".join(lines)
    return "\n".join(lines)


def _describe_xlsx(binary: bytes) -> str:
    workbook = load_workbook(BytesIO(binary), read_only=True, data_only=False)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets[:10]:
            lines.append(f"[工作表：{sheet.title}]")
            max_row = min(max(sheet.max_row, 1), 100)
            max_column = min(max(sheet.max_column, 1), 30)
            for row_index in range(1, max_row + 1):
                values = []
                for column_index in range(1, max_column + 1):
                    coordinate = f"{get_column_letter(column_index)}{row_index}"
                    value = sheet.cell(row=row_index, column=column_index).value
                    values.append(f"{coordinate}={value if value not in (None, '') else '<空>'}")
                lines.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(lines)


def _redact_descriptor(descriptor: str) -> str:
    descriptor = EMAIL_PATTERN.sub("[邮箱已脱敏]", descriptor)
    return PHONE_PATTERN.sub("[手机号已脱敏]", descriptor)


@dataclass(frozen=True)
class EditableSource:
    title: str
    binary_content: bytes | None


class AttachmentEditor:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def describe(self, user_id: UUID, attachment_id: UUID) -> tuple[str, str]:
        source = await self._owned_source(user_id, attachment_id)
        extension = Path(source.title).suffix.lower()
        if extension not in EDITABLE_EXTENSIONS:
            raise AppError(
                "ATTACHMENT_EDIT_UNSUPPORTED",
                "当前仅支持编辑 DOCX 和 XLSX",
                status_code=415,
            )
        if not source.binary_content:
            raise AppError(
                "ATTACHMENT_BINARY_MISSING",
                "该附件上传时未保存原文件，请重新上传后再编辑",
                status_code=422,
            )
        descriptor = (
            _describe_docx(source.binary_content)
            if extension == ".docx"
            else _describe_xlsx(source.binary_content)
        )
        return source.title, _redact_descriptor(descriptor)[:30_000]

    async def apply(
        self,
        *,
        user_id: UUID,
        conversation_id: UUID,
        attachment_id: UUID,
        user_request: str,
        changes: list[dict[str, str]],
        formatting: dict[str, str] | None = None,
        allow_unchanged: bool = False,
    ) -> tuple[dict, int]:
        source = await self._owned_source(user_id, attachment_id)
        if not source.binary_content:
            raise AppError("ATTACHMENT_BINARY_MISSING", "请重新上传原文件后再编辑", status_code=422)
        extension = Path(source.title).suffix.lower()
        if extension == ".docx":
            binary, applied = self._edit_docx(source.binary_content, changes)
            if formatting:
                binary, formatted = self._format_docx(binary, formatting)
                applied += formatted
            artifact_type, mime_type = "docx", DOCX_MIME
        elif extension == ".xlsx":
            binary, applied = self._edit_xlsx(source.binary_content, changes)
            artifact_type, mime_type = "xlsx", XLSX_MIME
        else:
            raise AppError(
                "ATTACHMENT_EDIT_UNSUPPORTED",
                "当前仅支持编辑 DOCX 和 XLSX",
                status_code=415,
            )
        if applied == 0 and not allow_unchanged:
            raise AppError(
                "ATTACHMENT_EDIT_NO_CHANGES",
                "没有找到可安全修改的位置，请更具体地说明要填写的内容",
                status_code=422,
            )
        artifact = await ArtifactService(self.session).create_binary(
            user_id=user_id,
            conversation_id=conversation_id,
            filename=_edited_filename(source.title, user_request),
            artifact_type=artifact_type,
            mime_type=mime_type,
            binary=binary,
        )
        return artifact, applied

    async def _owned_source(self, user_id: UUID, attachment_id: UUID) -> EditableSource:
        material = await self.session.scalar(
            select(WorkMaterialModel).where(
                WorkMaterialModel.id == str(attachment_id),
                WorkMaterialModel.user_id == str(user_id),
            )
        )
        if material is not None:
            return EditableSource(title=material.title, binary_content=material.binary_content)
        artifact = await self.session.scalar(
            select(GeneratedArtifactModel).where(
                GeneratedArtifactModel.id == str(attachment_id),
                GeneratedArtifactModel.user_id == str(user_id),
            )
        )
        if artifact is not None:
            return EditableSource(title=artifact.filename, binary_content=artifact.binary_content)
        raise AppError("ATTACHMENT_NOT_FOUND", "附件不存在或无权访问", status_code=404)

    @staticmethod
    def _edit_docx(binary: bytes, changes: list[dict[str, str]]) -> tuple[bytes, int]:
        document = Document(BytesIO(binary))
        applied = 0
        for change in changes:
            location = change.get("location", "")
            expected = change.get("expected_text", "")
            value = change["value"]
            table_match = DOCX_LOCATION.fullmatch(location)
            paragraph_match = DOCX_PARAGRAPH_LOCATION.fullmatch(location)
            try:
                if table_match:
                    table_index, row_index, cell_index = map(int, table_match.groups())
                    cell = document.tables[table_index].rows[row_index].cells[cell_index]
                    if _comparable_text(cell.text.strip()) != _comparable_text(expected):
                        continue
                    if _comparable_text(cell.text.strip()) == _comparable_text(value):
                        continue
                    _write_cell(cell, value)
                    applied += 1
                elif paragraph_match:
                    paragraph = document.paragraphs[int(paragraph_match.group(1))]
                    if _comparable_text(paragraph.text.strip()) != _comparable_text(expected):
                        continue
                    if _comparable_text(paragraph.text.strip()) == _comparable_text(value):
                        continue
                    _write_paragraph(paragraph, value)
                    applied += 1
            except (IndexError, ValueError):
                continue
        payload = BytesIO()
        document.save(payload)
        return payload.getvalue(), applied

    @staticmethod
    def _edit_xlsx(binary: bytes, changes: list[dict[str, str]]) -> tuple[bytes, int]:
        workbook = load_workbook(BytesIO(binary), data_only=False)
        applied = 0
        try:
            for change in changes:
                sheet_name = change.get("sheet", "")
                coordinate = change.get("cell", "").upper()
                if sheet_name not in workbook.sheetnames or not EXCEL_CELL.fullmatch(coordinate):
                    continue
                cell = workbook[sheet_name][coordinate]
                current = "" if cell.value in (None, "") else str(cell.value).strip()
                if current != change.get("expected_text", "") or current.startswith("="):
                    continue
                if current == change["value"].strip():
                    continue
                cell.value = change["value"]
                applied += 1
            payload = BytesIO()
            workbook.save(payload)
            return payload.getvalue(), applied
        finally:
            workbook.close()

    @staticmethod
    def _format_docx(binary: bytes, formatting: dict[str, str]) -> tuple[bytes, int]:
        document = Document(BytesIO(binary))
        target_color = formatting.get("font_color", "").upper()
        if not re.fullmatch(r"[0-9A-F]{6}", target_color):
            return binary, 0
        changed = 0
        roots = [document.element.body]
        seen_parts: set[str] = set()
        for section in document.sections:
            for part in (section.header.part, section.footer.part):
                if str(part.partname) in seen_parts:
                    continue
                seen_parts.add(str(part.partname))
                roots.append(part.element)
        for root in roots:
            for run in root.iter(qn("w:r")):
                if not any(node.text for node in run.iter(qn("w:t"))):
                    continue
                run_properties = run.get_or_add_rPr()
                color = run_properties.find(qn("w:color"))
                current = color.get(qn("w:val"), "").upper() if color is not None else ""
                has_theme = color is not None and any(
                    color.get(qn(attribute)) is not None
                    for attribute in ("w:themeColor", "w:themeTint", "w:themeShade")
                )
                if current == target_color and not has_theme:
                    continue
                if color is None:
                    color = OxmlElement("w:color")
                    run_properties.append(color)
                color.set(qn("w:val"), target_color)
                for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
                    color.attrib.pop(qn(attribute), None)
                changed += 1
        payload = BytesIO()
        document.save(payload)
        return payload.getvalue(), changed
