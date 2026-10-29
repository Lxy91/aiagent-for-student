"""Render Markdown in document order, including inline formatting and real links."""

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode


def parse_markdown(content: str) -> SyntaxTreeNode:
    return SyntaxTreeNode(MarkdownIt("commonmark", {"html": False}).enable("table").parse(content))


def _inline(paragraph, nodes, *, bold=False, italic=False, strike=False, target=None):
    for node in nodes:
        if node.type in {"strong", "em", "s"}:
            _inline(
                paragraph,
                node.children,
                bold=bold or node.type == "strong",
                italic=italic or node.type == "em",
                strike=strike or node.type == "s",
                target=target,
            )
        elif node.type == "link":
            href = node.attrGet("href") or ""
            if href.startswith(("https://", "http://", "mailto:")):
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
                paragraph._p.append(link)
                _inline(paragraph, node.children, bold=bold, italic=italic, target=link)
            else:
                _inline(paragraph, node.children, bold=bold, italic=italic, target=target)
        elif node.type in {"text", "code_inline", "softbreak", "hardbreak", "image"}:
            text = "\n" if node.type in {"softbreak", "hardbreak"} else node.content
            run = paragraph.add_run(text)
            run.bold, run.italic, run.font.strike = bold, italic, strike
            if node.type == "code_inline":
                run.font.name = "Consolas"
            if target is not None:
                target.append(run._r)
                run.font.color.rgb = RGBColor.from_string("185B8D")
                run.underline = True


def _numbering(document, ordered: bool, start: int) -> int:
    root = document.part.numbering_part.element
    abstract_id = (
        max(
            (int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))),
            default=-1,
        )
        + 1
    )
    num_id = max((int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))), default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    identifier = OxmlElement("w:nsid")
    identifier.set(qn("w:val"), f"{abstract_id + 65536:08X}")
    abstract.append(identifier)
    level_type = OxmlElement("w:multiLevelType")
    level_type.set(qn("w:val"), "singleLevel")
    abstract.append(level_type)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (
        ("start", str(start)),
        ("numFmt", "decimal" if ordered else "bullet"),
        ("lvlText", "%1." if ordered else "•"),
    ):
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:val"), value)
        level.append(element)
    abstract.append(level)
    first_num = root.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    restart = OxmlElement("w:startOverride")
    restart.set(qn("w:val"), str(start))
    override.append(restart)
    num.append(override)
    root.append(num)
    return num_id


def _render_table(document, node):
    rows = [n for section in node.children for n in section.children if n.type == "tr"]
    if not rows:
        return
    cols = len(rows[0].children)
    table = document.add_table(rows=0, cols=cols)
    table.autofit = False
    # Allocate more space to columns containing longer values.
    weights = [
        max(
            4,
            min(
                32,
                max(
                    len(row.children[i].children[0].content)
                    for row in rows
                    if i < len(row.children)
                ),
            ),
        )
        ** 0.6
        for i in range(cols)
    ]
    for col, weight in zip(table.columns, weights, strict=True):
        col.width = Mm(174 * weight / sum(weights))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn(f"w:{key}"), value)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for index, source_row in enumerate(rows):
        row = table.add_row()
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        # Keep ordinary rows together; long rows are allowed to paginate.
        if sum(len(cell.children[0].content) for cell in source_row.children) < 1000:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for col, source in enumerate(source_row.children):
            cell = row.cells[col]
            cell.width = table.columns[col].width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "70" if edge in {"top", "bottom"} else "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(
                qn("w:fill"), "E5EDF4" if index == 0 else ("F7F9FB" if index % 2 == 0 else "FFFFFF")
            )
            properties.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.1
            _inline(paragraph, source.children[0].children, bold=index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(10)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def render_docx(title: str, content: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(20)
    for name in ("Normal", "Title", "Subtitle", *[f"Heading {i}" for i in range(1, 7)]):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
        style.font.color.rgb = RGBColor(0, 0, 0)
        properties = style._element.get_or_add_pPr()
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
        style.paragraph_format.widow_control = True
        if name.startswith("Heading"):
            style.font.size = Pt(max(11, 18 - int(name[-1]) * 2))
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(14)
            style.paragraph_format.space_after = Pt(7)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(7)
    document.styles["Title"].font.size = Pt(24)
    document.styles["Title"].paragraph_format.keep_with_next = True
    document.add_paragraph(title, "Title")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.size = Pt(9)

    def render(nodes, *, depth=0, num_id=None, quote=False):
        for node in nodes:
            if node.type in {"bullet_list", "ordered_list"}:
                current_id = _numbering(
                    document, node.type == "ordered_list", int(node.attrGet("start") or 1)
                )
                render(node.children, depth=depth + 1, num_id=current_id, quote=quote)
            elif node.type == "list_item":
                first = True
                for child in node.children:
                    render([child], depth=depth, num_id=num_id if first else None, quote=quote)
                    first = False
            elif node.type == "blockquote":
                render(node.children, depth=depth, quote=True)
            elif node.type == "table":
                _render_table(document, node)
            elif node.type in {"paragraph", "heading"}:
                style = f"Heading {node.tag[1:]}" if node.type == "heading" else "Normal"
                paragraph = document.add_paragraph(style=style)
                if depth or quote:
                    paragraph.paragraph_format.left_indent = Mm(6 * max(depth, 1))
                if num_id is not None and node.type == "paragraph":
                    properties = paragraph._p.get_or_add_pPr()
                    numbering = OxmlElement("w:numPr")
                    for tag, value in (("ilvl", "0"), ("numId", str(num_id))):
                        item = OxmlElement(f"w:{tag}")
                        item.set(qn("w:val"), value)
                        numbering.append(item)
                    properties.append(numbering)
                _inline(paragraph, node.children[0].children)
            elif node.type in {"fence", "code_block"}:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(node.content.rstrip())
                run.font.name, run.font.size = "Consolas", Pt(9)

    render(parse_markdown(content).children)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
