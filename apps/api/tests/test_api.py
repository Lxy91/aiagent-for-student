from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.infrastructure.search.tavily import TavilyWebSearchProvider
from app.modules.artifacts.editor import (
    AttachmentEditor,
    _describe_docx,
    allows_unfilled_fields,
    editable_attachments,
    normalize_edit_plan,
    requested_docx_formatting,
    select_edit_target,
)
from app.modules.artifacts.service import (
    artifact_generation_instruction,
    build_docx,
    build_xlsx,
    requested_artifact_types,
)
from app.modules.chat.api import (
    artifact_execution_narrative,
    requires_image_generation,
    requires_web_search,
)
from app.modules.growth.service import GrowthService
from app.ports.web_search import SearchResult, WebSearchResponse


def build_text_pdf(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    payload = BytesIO()
    writer.write(payload)
    return payload.getvalue()


def test_health_and_ready(client: TestClient) -> None:
    assert client.get("/api/v1/health").json()["status"] == "ok"
    ready = client.get("/api/v1/ready")
    assert ready.status_code == 200
    assert ready.json()["dependencies"]["deepseek"] == "demo-mode"


def test_web_search_routing_covers_current_information_without_overmatching() -> None:
    assert requires_web_search("总结一下今日新闻")
    assert requires_web_search("最近有哪些校招政策变化？")
    assert requires_web_search("帮我查一下这家公司")
    assert not requires_web_search("我当前很焦虑，该怎么和领导沟通？")
    assert not requires_web_search("帮我准备第一次周会")


def test_image_generation_routing_supports_explicit_and_graphic_text_requests() -> None:
    assert requires_image_generation("生成一张图：第一次参加周会")
    assert requires_image_generation("生成一张女娲五杀图片")
    assert requires_image_generation("帮我画一个大学生第一次参加周会的场景")
    assert requires_image_generation("制作一份秋招主题海报")
    assert requires_image_generation("给我来张职场风格插画")
    assert requires_image_generation("把这段内容做成图文版")
    assert requires_image_generation("请给这份建议配一张图")
    assert not requires_image_generation("分析这张图片里的问题")
    assert not requires_image_generation("如何生成一张图片？")
    assert not requires_image_generation("生成图片的方法有哪些？")
    assert not requires_image_generation("这个系统是否支持图片生成功能？")


def test_artifact_routing_only_matches_explicit_download_requests() -> None:
    assert requested_artifact_types("请生成Word文档和Excel表格") == ("docx", "xlsx")
    assert requested_artifact_types("你给我生成份word文档") == ("docx",)
    assert requested_artifact_types("根据这个面经写一份word报告") == ("docx",)
    assert requested_artifact_types("把材料整理成一个Word报告") == ("docx",)
    assert requested_artifact_types("把回答导出文档") == ("docx",)
    assert requested_artifact_types("请分析我上传的文档和Excel") == ()
    assert requested_artifact_types("如何写word报告") == ()
    assert requested_artifact_types("分析这份Word报告") == ()
    instruction = artifact_generation_instruction("你给我生成份word文档")
    assert "具备生成并提供可下载的 Word（DOCX）文件的能力" in instruction
    assert "不得声称无法生成" in instruction
    assert artifact_generation_instruction("帮我分析一下职业规划") == ""


def test_generated_artifact_uses_ai_authored_document_title() -> None:
    prompt = "这是什么，根据这个帮我写一份自我介绍并生成 Word 文档"
    content = "# 前端岗位面试自我介绍稿\n\n## 基本信息\n我是一名前端开发实习生。"

    docx_filename, _ = build_docx(prompt, content)
    xlsx_filename, _ = build_xlsx(prompt, content)

    assert docx_filename == "前端岗位面试自我介绍稿.docx"
    assert xlsx_filename == "前端岗位面试自我介绍稿.xlsx"
    assert "这是什么" not in docx_filename
    narrative = artifact_execution_narrative(("docx",), [])
    assert narrative.understanding == "你希望把本次内容整理成Word 文档并直接下载。"
    assert narrative.next_action == "我会整理适合写入文件的完整内容，并生成可下载的Word 文档。"


def test_attachment_edit_routing_requires_edit_intent_and_editable_file() -> None:
    attachments = [
        {
            "id": "document-id",
            "title": "实习问卷.docx",
            "material_type": "document",
            "mime_type": "application/octet-stream",
            "status": "ready",
        }
    ]
    assert editable_attachments("这个补全发给我", attachments) == attachments
    assert editable_attachments("总结一下这个文件", attachments) == []
    assert editable_attachments(
        "修改一下这张图片",
        [{**attachments[0], "title": "screenshot.png", "material_type": "image"}],
    ) == []
    assert allows_unfilled_fields("不清楚的地方可以先不填写")
    assert allows_unfilled_fields("不知道的内容留空即可")
    assert not allows_unfilled_fields("缺少内容，请先问我")
    assert requested_docx_formatting("把字体颜色统一一下", "报告.docx") == {
        "font_color": "000000"
    }
    assert requested_docx_formatting("把字体颜色统一一下", "报告.xlsx") == {}


def test_attachment_edit_selects_numbered_target_and_uses_other_docx_as_reference() -> None:
    attachments = [
        {"id": "target", "title": "问卷.docx"},
        {"id": "reference", "title": "实习内容.docx"},
    ]
    assert select_edit_target("根据文档2的内容补全文档1", attachments) == attachments[0]
    assert select_edit_target("用文件1的内容修改文件2", attachments) == attachments[1]
    assert select_edit_target("补全这两个文件", attachments) is None


def test_docx_edit_accepts_model_compressed_template_whitespace() -> None:
    template = Document()
    table = template.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "评价：\n\n\n\n签名：        年   月   日"
    payload = BytesIO()
    template.save(payload)

    edited_binary, applied = AttachmentEditor._edit_docx(
        payload.getvalue(),
        [
            {
                "location": "table:0:row:0:cell:0",
                "expected_text": "评价：\n\n签名： 年 月 日",
                "value": "评价：表现优秀\n\n签名：        年   月   日",
            }
        ],
    )

    assert applied == 1
    edited = Document(BytesIO(edited_binary))
    assert edited.tables[0].cell(0, 0).text == "评价：表现优秀\n\n签名：        年   月   日"


def test_docx_descriptor_lists_merged_cell_only_once() -> None:
    template = Document()
    table = template.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 2))
    merged.text = "合并填写区域"
    payload = BytesIO()
    template.save(payload)

    descriptor = _describe_docx(payload.getvalue())

    assert descriptor.count("合并填写区域") == 1
    assert "table:0:row:0:cell:0" in descriptor


def test_docx_edit_does_not_count_unchanged_values_as_modifications() -> None:
    template = Document()
    table = template.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "已经填写"
    payload = BytesIO()
    template.save(payload)

    edited_binary, applied = AttachmentEditor._edit_docx(
        payload.getvalue(),
        [
            {
                "location": "table:0:row:0:cell:0",
                "expected_text": "已经填写",
                "value": "已经填写",
            }
        ],
    )

    assert applied == 0
    assert Document(BytesIO(edited_binary)).tables[0].cell(0, 0).text == "已经填写"


def test_docx_descriptor_escapes_multiline_cell_as_one_json_string() -> None:
    template = Document()
    table = template.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "标题\n\n签名："
    payload = BytesIO()
    template.save(payload)

    descriptor = _describe_docx(payload.getvalue())

    assert descriptor.count("table:0:row:0:cell:0") == 1
    assert '"标题\\n\\n签名："' in descriptor


def test_edit_plan_never_writes_blank_marker_as_cell_content() -> None:
    _, _, changes = normalize_edit_plan(
        {
            "changes": [
                {"location": "table:0:row:0:cell:0", "expected_text": "<空>", "value": "<空>"},
                {"location": "table:0:row:1:cell:0", "expected_text": "<空>", "value": "已填写"},
            ]
        }
    )

    assert changes == [
        {
            "location": "table:0:row:1:cell:0",
            "sheet": "",
            "cell": "",
            "expected_text": "",
            "value": "已填写",
        }
    ]


def test_docx_edit_preserves_form_heading_signature_and_response_space() -> None:
    template = Document()
    table = template.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "评价与建议："
    for _ in range(4):
        cell.add_paragraph("")
    cell.add_paragraph("负责人签名：        年   月   日")
    payload = BytesIO()
    template.save(payload)
    expected = cell.text.strip()

    edited_binary, applied = AttachmentEditor._edit_docx(
        payload.getvalue(),
        [
            {
                "location": "table:0:row:0:cell:0",
                "expected_text": expected,
                "value": (
                    "评价与建议：\n该同学实习期间表现认真，能够按时完成任务。\n"
                    "负责人签名：        年   月   日"
                ),
            }
        ],
    )

    edited_cell = Document(BytesIO(edited_binary)).tables[0].cell(0, 0)
    assert applied == 1
    assert edited_cell.paragraphs[0].text == "评价与建议："
    assert edited_cell.paragraphs[1].text == "该同学实习期间表现认真，能够按时完成任务。"
    assert edited_cell.paragraphs[-1].text == "负责人签名：        年   月   日"


def test_auth_rejects_duplicate_email(client: TestClient, auth_headers: dict[str, str]) -> None:
    del auth_headers
    response = client.post(
        "/api/v1/auth/register",
        json={
            "email": "student@example.com",
            "password": "another-password",
            "display_name": "另一位同学",
        },
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "EMAIL_ALREADY_EXISTS"
    assert response.json()["error"]["trace_id"]


def test_conversation_stream(client: TestClient, auth_headers: dict[str, str]) -> None:
    created = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "任务澄清", "mode": "standard"},
    )
    assert created.status_code == 201
    conversation_id = created.json()["id"]

    response = client.post(
        f"/api/v1/conversations/{conversation_id}/messages:stream",
        headers=auth_headers,
        json={"content": "如何向领导澄清需求？"},
    )
    assert response.status_code == 200
    assert "event: message.started" in response.text
    assert "event: message.delta" in response.text
    assert "event: message.completed" in response.text
    assert '"demo_mode": true' in response.text

    messages = client.get(f"/api/v1/conversations/{conversation_id}/messages", headers=auth_headers)
    assert messages.status_code == 200
    assert [item["role"] for item in messages.json()] == ["user", "assistant"]


def test_conversation_rename_pin_and_archive(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    first = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "第一个会话", "mode": "standard"},
    ).json()
    second = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "第二个会话", "mode": "standard"},
    ).json()

    updated = client.patch(
        f"/api/v1/conversations/{first['id']}",
        headers=auth_headers,
        json={"title": "  重命名后的会话  ", "is_pinned": True},
    )
    assert updated.status_code == 200
    assert updated.json()["title"] == "重命名后的会话"
    assert updated.json()["is_pinned"] is True

    conversations = client.get("/api/v1/conversations", headers=auth_headers).json()
    assert conversations[0]["id"] == first["id"]

    archived = client.patch(
        f"/api/v1/conversations/{first['id']}",
        headers=auth_headers,
        json={"is_archived": True},
    )
    assert archived.status_code == 200
    conversations = client.get("/api/v1/conversations", headers=auth_headers).json()
    assert [item["id"] for item in conversations] == [second["id"]]


def test_first_question_becomes_default_conversation_title(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "新对话", "mode": "standard"},
    ).json()

    with client.stream(
        "POST",
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "请帮我准备第一次周会？"},
    ) as response:
        assert response.status_code == 200
        list(response.iter_lines())

    updated = client.get("/api/v1/conversations", headers=auth_headers).json()[0]
    assert updated["title"] == "准备第一次周会"


def test_first_question_does_not_overwrite_manual_conversation_title(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "我的自定义名称", "mode": "standard"},
    ).json()

    with client.stream(
        "POST",
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "这是第一个问题"},
    ) as response:
        assert response.status_code == 200
        list(response.iter_lines())

    updated = client.get("/api/v1/conversations", headers=auth_headers).json()[0]
    assert updated["title"] == "我的自定义名称"


def test_memory_confirmation_and_delete(client: TestClient, auth_headers: dict[str, str]) -> None:
    candidate = client.post(
        "/api/v1/memory-candidates",
        headers=auth_headers,
        json={"type": "goal", "content": "三个月内独立完成需求评审"},
    )
    assert candidate.status_code == 201
    memory = client.post(
        f"/api/v1/memory-candidates/{candidate.json()['id']}:confirm",
        headers=auth_headers,
        json={},
    )
    assert memory.status_code == 201
    memory_id = memory.json()["id"]
    assert len(client.get("/api/v1/memories", headers=auth_headers).json()) == 1
    assert client.delete(f"/api/v1/memories/{memory_id}", headers=auth_headers).status_code == 204


def test_knowledge_upload(client: TestClient, auth_headers: dict[str, str]) -> None:
    response = client.post(
        "/api/v1/knowledge/documents",
        headers=auth_headers,
        files={"file": ("guide.md", b"# Workplace guide\nAsk clear questions.", "text/markdown")},
        data={"source": "测试资料"},
    )
    assert response.status_code == 202
    assert response.json()["status"] == "ready"
    assert response.json()["source"] == "测试资料"


def test_plan_confirmation_is_idempotent(client: TestClient, auth_headers: dict[str, str]) -> None:
    draft = client.post(
        "/api/v1/plans:draft",
        headers=auth_headers,
        json={"goal": "三天看懂 React 项目并完成汇报", "deadline": "2026-08-23"},
    )
    assert draft.status_code == 201
    body = draft.json()
    confirmation = {"confirmation_token": body["confirmation_token"]}
    first = client.post(
        f"/api/v1/plan-drafts/{body['id']}:confirm", headers=auth_headers, json=confirmation
    )
    second = client.post(
        f"/api/v1/plan-drafts/{body['id']}:confirm", headers=auth_headers, json=confirmation
    )
    assert first.status_code == 201
    assert second.status_code == 201
    assert first.json()["id"] == second.json()["id"]

    task_id = first.json()["tasks"][0]["id"]
    updated = client.patch(
        f"/api/v1/tasks/{task_id}", headers=auth_headers, json={"status": "done"}
    )
    assert updated.status_code == 200
    assert updated.json()["status"] == "done"


def test_tool_catalog_and_calendar_draft_are_idempotent(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    catalog = client.get("/api/v1/tools", headers=auth_headers)
    assert catalog.status_code == 200
    tools = {item["name"]: item for item in catalog.json()}
    assert tools["web.search"]["effect"] == "read"
    assert tools["calendar.create_event_draft"]["effect"] == "draft"

    payload = {
        "arguments": {
            "title": "准备周会",
            "start_at": "2026-08-24T09:00:00+08:00",
            "end_at": "2026-08-24T09:30:00+08:00",
            "timezone": "Asia/Shanghai",
            "notes": "检查本周进展",
        },
        "idempotency_key": "calendar-test-001",
    }
    first = client.post(
        "/api/v1/tools/calendar.create_event_draft:invoke",
        headers=auth_headers,
        json=payload,
    )
    second = client.post(
        "/api/v1/tools/calendar.create_event_draft:invoke",
        headers=auth_headers,
        json=payload,
    )
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["result"]["action_id"] == second.json()["result"]["action_id"]
    assert first.json()["result"]["status"] == "draft"


def test_web_search_tool_returns_sources_and_audit(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    async def fake_search(self, query: str, *, max_results: int) -> WebSearchResponse:
        del self, max_results
        return WebSearchResponse(
            query=query,
            results=[
                SearchResult(
                    title="教育部就业服务平台",
                    url="https://www.ncss.cn/example",
                    snippet="公开招聘与就业服务信息。",
                    source="www.ncss.cn",
                    score=0.98,
                )
            ],
        )

    monkeypatch.setattr(TavilyWebSearchProvider, "search", fake_search)
    response = client.post(
        "/api/v1/tools/web.search:invoke",
        headers=auth_headers,
        json={"arguments": {"query": "2026 校招最新信息", "max_results": 3}},
    )
    assert response.status_code == 200
    assert response.json()["result"]["results"][0]["url"] == "https://www.ncss.cn/example"

    runs = client.get("/api/v1/tools/runs", headers=auth_headers)
    assert runs.status_code == 200
    assert runs.json()[0]["tool_name"] == "web.search"
    assert runs.json()[0]["status"] == "succeeded"
    assert runs.json()[0]["input_summary"]["query"] == "2026 校招最新信息"


def test_chat_stream_emits_tool_events_and_persists_citations(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    selections = 0

    async def fake_select(self, messages, tools, *, force_tool_name=None, allow_tool_calls=True):
        nonlocal selections
        del self, messages, tools, force_tool_name
        if not allow_tool_calls or selections:
            return {
                "role": "assistant",
                "content": "根据检索结果，建议优先查看官方就业信息 [1]。",
            }, []
        selections += 1
        call = {
            "id": "call-search-1",
            "type": "function",
            "function": {
                "name": "web__search",
                "arguments": '{"query":"2026 校招最新政策","max_results":3}',
            },
        }
        return {"role": "assistant", "content": None, "tool_calls": [call]}, [call]

    async def fake_search(self, query: str, *, max_results: int) -> WebSearchResponse:
        del self, max_results
        return WebSearchResponse(
            query=query,
            results=[
                SearchResult(
                    title="官方就业信息",
                    url="https://example.edu.cn/jobs",
                    snippet="最新就业政策摘要。",
                    source="example.edu.cn",
                )
            ],
        )

    async def fake_narrative(self, messages, *, max_tokens=2400):
        del self, max_tokens
        no_tool = "不调用工具，直接回答" in messages[-1]["content"]
        return {
            "understanding": "用户想了解 2026 年校招政策的最新变化和实际影响。",
            "next_action": (
                "我会结合已经取得的就业信息整理政策要点。"
                if no_tool
                else "我会联网检索权威就业来源，再根据结果整理政策要点。"
            ),
        }

    async def fake_stream(self, messages):
        del self, messages
        for chunk in ("根据检索结果，", "建议优先查看", "官方就业信息 [1]。"):
            yield chunk

    monkeypatch.setattr(DeepSeekProvider, "select_tool_calls", fake_select)
    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_narrative)
    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)
    monkeypatch.setattr(TavilyWebSearchProvider, "search", fake_search)
    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(Settings, "online_search_available", property(lambda self: True))

    created = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "联网检索", "mode": "standard"},
    )
    conversation_id = created.json()["id"]
    response = client.post(
        f"/api/v1/conversations/{conversation_id}/messages:stream",
        headers=auth_headers,
        json={"content": "2026 校招最新政策有哪些？"},
    )
    assert response.status_code == 200
    assert "event: tool.started" in response.text
    assert "event: tool.completed" in response.text
    assert "event: reasoning.step" in response.text
    assert "event: citations" in response.text
    assert "https://example.edu.cn/jobs" in response.text
    assert "校招政策的最新变化" in response.text
    assert "目标、时效性要求和当前可用工具" not in response.text
    assert "DSML" not in response.text

    messages = client.get(
        f"/api/v1/conversations/{conversation_id}/messages", headers=auth_headers
    ).json()
    assert messages[-1]["citations"][0]["source"] == "example.edu.cn"
    assert messages[-1]["reasoning_steps"][-1]["status"] == "completed"


def test_chat_trace_is_contextual_when_no_tool_is_selected(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    async def fake_select(self, messages, tools, *, force_tool_name=None, allow_tool_calls=True):
        del self, messages, tools, force_tool_name, allow_tool_calls
        return {
            "role": "assistant",
            "content": "我只能依据当前会话中你主动提供的信息来认识你。",
        }, []

    async def fake_narrative(self, messages, *, max_tokens=2400):
        del self, messages, max_tokens
        return {
            "understanding": "你想确认助手目前能识别哪些关于你的信息。",
            "next_action": "我会依据当前会话说明我能看到的信息，以及无法判断的部分。",
        }

    async def fake_stream(self, messages):
        del self, messages
        for chunk in ("我只能依据当前会话中", "你主动提供的信息", "来认识你。"):
            yield chunk

    monkeypatch.setattr(DeepSeekProvider, "select_tool_calls", fake_select)
    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_narrative)
    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)
    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))

    created = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "身份问题", "mode": "standard"},
    )
    conversation_id = created.json()["id"]
    response = client.post(
        f"/api/v1/conversations/{conversation_id}/messages:stream",
        headers=auth_headers,
        json={"content": "我是谁"},
    )

    assert response.status_code == 200
    assert "目前能识别哪些关于你的信息" in response.text
    assert "依据当前会话说明" in response.text
    assert "时效性要求和当前可用工具" not in response.text
    assert "运行了 web.search" not in response.text
    assert response.text.count("event: message.delta") == 3


def test_chat_trace_uses_parsed_docx_attachment_state_instead_of_model_guess(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    final_messages: list[dict] = []

    async def fake_select(self, messages, tools, *, force_tool_name=None, allow_tool_calls=True):
        del self, messages, tools, force_tool_name, allow_tool_calls
        return {"role": "assistant", "content": "我会总结文档。"}, []

    async def wrong_narrative(self, messages, *, max_tokens=2400):
        del self, messages, max_tokens
        return {
            "understanding": "用户没有提供任何文档内容，无法总结。",
            "next_action": "我会请用户重新提供文档。",
        }

    async def fake_stream(self, messages):
        del self
        final_messages.extend(messages)
        yield "这是一份实习单位问卷调查表。"

    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(DeepSeekProvider, "select_tool_calls", fake_select)
    monkeypatch.setattr(DeepSeekProvider, "complete_json", wrong_narrative)
    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)

    payload = BytesIO()
    with ZipFile(payload, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>实习单位问卷调查表，"
            "由企业填写并加盖公章。</w:t></w:r></w:p></w:body></w:document>",
        )
    material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "实习单位问卷调查表.docx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"purpose": "conversation"},
    ).json()
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "文档总结", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "总结一下这个文档", "attachment_ids": [material["id"]]},
    )

    assert response.status_code == 200
    assert "已收到 1 个附件" in response.text
    assert "实习单位问卷调查表.docx（已解析文字内容）" in response.text
    assert "没有提供任何文档内容" not in response.text
    assert "重新提供文档" not in response.text
    assert "实习单位问卷调查表，由企业填写并加盖公章" in final_messages[-1]["content"]


def test_docx_extraction_reads_text_boxes_without_including_ooxml_markup() -> None:
    payload = BytesIO()
    with ZipFile(payload, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            "<w:body><w:p><w:r><w:t>简历标题</w:t></w:r></w:p>"
            "<w:p><w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
            "<w:p><w:r><w:t>文本框中的工作经历</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
            "</w:body></w:document>",
        )

    extracted = GrowthService._extract_text(".docx", payload.getvalue())

    assert extracted == "简历标题\n文本框中的工作经历"
    assert "<w:" not in extracted


def test_corrupt_docx_returns_explicit_parse_error(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "broken.docx",
                b"not-a-zip-file",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"purpose": "conversation"},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "DOCX_PARSE_FAILED"


def test_ai_planner_uses_structured_model_breakdown(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    async def fake_complete_json(self, messages, *, max_tokens=2400):
        del self, messages, max_tokens
        return {
            "title": "完成首次周会汇报",
            "deliverable": "一份可讲解的周会材料",
            "assumptions": ["每天可投入一小时"],
            "risks": ["数据收集不完整"],
            "tasks": [
                {
                    "title": "收集本周事实",
                    "description": "整理完成事项、问题和数据。",
                    "estimated_minutes": 45,
                    "priority": "high",
                    "due_offset_days": 0,
                    "done_definition": "形成事实清单。",
                    "depends_on_indexes": [],
                },
                {
                    "title": "形成汇报材料",
                    "description": "按结论、进展、风险和下一步组织内容。",
                    "estimated_minutes": 60,
                    "priority": "high",
                    "due_offset_days": 1,
                    "done_definition": "完成可预演的材料。",
                    "depends_on_indexes": [0],
                },
            ],
        }

    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_complete_json)
    response = client.post(
        "/api/v1/plans:draft",
        headers=auth_headers,
        json={"goal": "准备第一次周会并完成清晰汇报", "deadline": "2026-08-25"},
    )
    assert response.status_code == 201
    assert [task["title"] for task in response.json()["tasks"]] == [
        "收集本周事实",
        "形成汇报材料",
    ]
    assert response.json()["tasks"][1]["depends_on"] == [response.json()["tasks"][0]["id"]]


def test_explicit_search_never_falls_back_to_model_when_unconfigured(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings

    monkeypatch.setattr(Settings, "online_search_available", property(lambda self: False))
    created = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "搜索降级", "mode": "standard"},
    )
    response = client.post(
        f"/api/v1/conversations/{created.json()['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "请联网搜索今天的校招信息"},
    )
    assert response.status_code == 200
    assert "event: error" in response.text
    assert "TOOL_UNAVAILABLE" in response.text
    assert "event: message.delta" not in response.text


def test_v03_meeting_material_creates_traceable_minutes_and_redacts_privacy(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    content = (
        "需求评审会\n"
        "本周确定新人引导流程。\n"
        "行动项：小林负责周五前完成需求说明。\n"
        "待确认：演示时间是否调整到周六？\n"
        "联系邮箱 intern@example.com，手机 13800138000。"
    ).encode()
    response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("meeting.md", content, "text/markdown")},
        data={"purpose": "meeting"},
    )
    assert response.status_code == 202
    material = response.json()
    assert material["status"] == "ready"
    assert material["privacy_status"] == "redacted"
    assert "intern@example.com" not in material["content_excerpt"]
    assert "13800138000" not in material["content_excerpt"]
    assert material["minutes"]["action_items"] == ["行动项：小林负责周五前完成需求说明。"]
    assert "演示时间" in material["minutes"]["pending_facts"][0]

    profile = client.get("/api/v1/growth-profile", headers=auth_headers).json()
    assert profile["evidence"][0]["source_id"] == material["id"]
    assert profile["evidence"][0]["source_title"] == "meeting.md"


def test_v03_non_text_multimodal_input_does_not_invent_content(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("meeting.png", b"not-a-real-image", "image/png")},
        data={"purpose": "meeting"},
    )
    assert response.status_code == 202
    material = response.json()
    assert material["status"] == "needs_confirmation"
    assert material["minutes"]["action_items"] == []
    assert "未配置音频/图像识别" in material["minutes"]["summary"]


def test_v03_chat_accepts_multiple_attachment_types_as_grounded_context(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings
    from app.infrastructure.llm.bigmodel import BigModelProvider
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    captured_messages: list[dict] = []
    captured_images: list[dict] = []

    async def fake_select(self, messages, tools, *, force_tool_name=None, allow_tool_calls=True):
        del self, messages, tools, force_tool_name, allow_tool_calls
        return {"role": "assistant", "content": "我会依据附件内容回答。"}, []

    async def fake_narrative(self, messages, *, max_tokens=2400):
        del self, messages, max_tokens
        return {
            "understanding": "用户希望结合已上传的附件进行总结。",
            "next_action": "我会读取可用文字，并明确无法识别的附件。",
        }

    async def fake_vision(self, messages, images):
        del self
        captured_messages.extend(messages)
        captured_images.extend(images)
        yield "已根据图片和可读取的附件内容完成总结。"

    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(Settings, "bigmodel_enabled", property(lambda self: True))
    monkeypatch.setattr(DeepSeekProvider, "select_tool_calls", fake_select)
    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_narrative)
    monkeypatch.setattr(BigModelProvider, "stream_vision", fake_vision)

    text_material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "brief.txt",
                "本周完成需求评审，联系 intern@example.com".encode(),
                "text/plain",
            )
        },
        data={"purpose": "conversation"},
    ).json()
    image_material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("board.png", b"not-a-real-image", "image/png")},
        data={"purpose": "conversation"},
    ).json()
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "新对话", "mode": "standard"},
    ).json()

    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={
            "content": "请总结附件",
            "attachment_ids": [text_material["id"], image_material["id"]],
        },
    )
    assert response.status_code == 200
    assert "event: message.completed" in response.text
    model_context = captured_messages[-1]["content"]
    assert "本周完成需求评审" in model_context
    assert "intern@example.com" not in model_context
    assert "[邮箱已脱敏]" in model_context
    assert "图片内容已作为视觉输入提供" in model_context
    assert captured_images[0]["title"] == "board.png"
    assert captured_images[0]["base64"]

    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    assert [item["title"] for item in messages[0]["attachments"]] == [
        "brief.txt",
        "board.png",
    ]
    assert messages[0]["content"] == "请总结附件"


def test_v03_chat_routes_explicit_image_generation_to_glm_image(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings
    from app.infrastructure.llm.bigmodel import BigModelProvider

    async def fake_generate(self, prompt: str, size: str = "1280x1280") -> list[str]:
        del self, size
        assert "生成一张图" in prompt
        return ["https://example.com/generated.png"]

    monkeypatch.setattr(Settings, "bigmodel_enabled", property(lambda self: True))
    monkeypatch.setattr(BigModelProvider, "generate_image", fake_generate)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "图片生成", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "生成一张图：大学生第一次参加周会"},
    )
    assert response.status_code == 200
    assert "event: images" in response.text
    assert "https://example.com/generated.png" in response.text
    assert '"model": "glm-image"' in response.text
    assert "已为你生成图文内容" in response.text
    assert "图片主题" in response.text

    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    assert messages[-1]["generated_images"][0]["url"] == "https://example.com/generated.png"


def test_v03_excel_and_csv_materials_are_read_and_redacted(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "本周任务"
    sheet.append(["事项", "负责人", "联系方式"])
    sheet.append(["行动项：完成竞品分析", "小林", "intern@example.com"])
    feedback = workbook.create_sheet("反馈")
    feedback.append(["日期", "反馈"])
    feedback.append(["2026-08-24", "待确认：汇报时间是否调整？"])
    payload = BytesIO()
    workbook.save(payload)

    excel_response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "weekly.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"purpose": "meeting"},
    )
    assert excel_response.status_code == 202
    excel = excel_response.json()
    assert excel["material_type"] == "spreadsheet"
    assert excel["status"] == "ready"
    assert "[工作表：本周任务]" in excel["content_excerpt"]
    assert "[工作表：反馈]" in excel["content_excerpt"]
    assert "intern@example.com" not in excel["content_excerpt"]
    assert "[邮箱已脱敏]" in excel["content_excerpt"]
    assert any("完成竞品分析" in item for item in excel["minutes"]["action_items"])
    assert any("汇报时间" in item for item in excel["minutes"]["pending_facts"])

    csv_response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "tasks.csv",
                "事项,状态\n整理周报,已完成\n".encode("gb18030"),
                "text/csv",
            )
        },
        data={"purpose": "reference"},
    )
    assert csv_response.status_code == 202
    assert csv_response.json()["material_type"] == "spreadsheet"
    assert "表头：事项 | 状态" in csv_response.json()["content_excerpt"]


def test_v03_pdf_material_is_extracted_for_chat_context(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.core.config import Settings
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    captured_messages: list[dict] = []

    async def fake_select(self, messages, tools, *, force_tool_name=None, allow_tool_calls=True):
        del self, messages, tools, force_tool_name, allow_tool_calls
        return {"role": "assistant", "content": "I will summarize the PDF."}, []

    async def fake_narrative(self, messages, *, max_tokens=2400):
        del self, messages, max_tokens
        return {
            "understanding": "The user wants a PDF summary.",
            "next_action": "I will use the extracted PDF text.",
        }

    async def fake_stream(self, messages):
        del self
        captured_messages.extend(messages)
        yield "PDF summary completed."

    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(DeepSeekProvider, "select_tool_calls", fake_select)
    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_narrative)
    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)

    material_response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "weekly-review.pdf",
                build_text_pdf("Weekly review action item due Friday"),
                "application/pdf",
            )
        },
        data={"purpose": "conversation"},
    )
    assert material_response.status_code == 202
    material = material_response.json()
    assert material["status"] == "ready"
    assert "[PDF 第 1 页]" in material["content_excerpt"]
    assert "Weekly review action item due Friday" in material["content_excerpt"]

    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "PDF summary", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "Summarize the PDF", "attachment_ids": [material["id"]]},
    )
    assert response.status_code == 200
    assert "event: message.completed" in response.text
    assert "Weekly review action item due Friday" in captured_messages[-1]["content"]


def test_v03_rejects_broken_pdf_instead_of_silently_accepting_it(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("broken.pdf", b"not-a-pdf-file", "application/pdf")},
        data={"purpose": "conversation"},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "PDF_PARSE_FAILED"


def test_v03_rejects_broken_excel_instead_of_guessing(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    response = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("broken.xlsx", b"not-an-excel-file", "application/octet-stream")},
        data={"purpose": "meeting"},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "SPREADSHEET_PARSE_FAILED"


def test_v03_report_uses_completed_tasks_and_deletes_material_derivatives(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    draft = client.post(
        "/api/v1/plans:draft",
        headers=auth_headers,
        json={"goal": "完成本周汇报", "deadline": "2026-08-24"},
    ).json()
    plan = client.post(
        f"/api/v1/plan-drafts/{draft['id']}:confirm",
        headers=auth_headers,
        json={"confirmation_token": draft["confirmation_token"]},
    ).json()
    first_task = plan["tasks"][0]
    client.patch(
        f"/api/v1/tasks/{first_task['id']}",
        headers=auth_headers,
        json={"status": "done"},
    )
    material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("review.txt", b"weekly review notes", "text/plain")},
        data={"purpose": "meeting"},
    ).json()
    report_response = client.post(
        "/api/v1/reports:draft",
        headers=auth_headers,
        json={
            "period_type": "weekly",
            "period_start": "2026-08-01",
            "period_end": "2026-08-31",
        },
    )
    assert report_response.status_code == 201
    report = report_response.json()
    source_ids = {source["id"] for source in report["sources"]}
    assert first_task["id"] in source_ids
    assert material["id"] in source_ids
    assert first_task["title"] in report["sections"]["completed"][0]

    deleted = client.delete(f"/api/v1/materials/{material['id']}", headers=auth_headers)
    assert deleted.status_code == 204
    missing = client.get(f"/api/v1/materials/{material['id']}", headers=auth_headers)
    assert missing.status_code == 404
    assert client.get("/api/v1/reports", headers=auth_headers).json() == []
    evidence = client.get("/api/v1/growth-profile", headers=auth_headers).json()["evidence"]
    assert all(item["source_id"] != material["id"] for item in evidence)


def test_chat_generated_docx_and_xlsx_are_downloadable_and_private(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    async def fake_stream(self, messages):
        del self, messages
        yield "# 周会准备\n\n以下是可执行安排：\n\n"
        yield "| 事项 | 负责人 | 截止时间 |\n| --- | --- | --- |\n"
        yield "| 整理进展 | 小林 | 周五 |\n| 确认风险 | 小周 | 周四 |"

    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "文件导出", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "请生成Word文档和Excel表格：周会准备"},
    )
    assert response.status_code == 200
    assert "event: artifacts" in response.text

    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    artifacts = messages[-1]["generated_artifacts"]
    assert {item["artifact_type"] for item in artifacts} == {"docx", "xlsx"}

    by_type = {item["artifact_type"]: item for item in artifacts}
    docx_response = client.get(
        f"/api/v1/artifacts/{by_type['docx']['id']}/download", headers=auth_headers
    )
    assert docx_response.status_code == 200
    assert docx_response.content.startswith(b"PK")
    assert "filename*=UTF-8''" in docx_response.headers["content-disposition"]
    document = Document(BytesIO(docx_response.content))
    assert "周会准备" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.tables[0].cell(1, 0).text == "整理进展"

    xlsx_response = client.get(
        f"/api/v1/artifacts/{by_type['xlsx']['id']}/download", headers=auth_headers
    )
    assert xlsx_response.status_code == 200
    workbook = load_workbook(BytesIO(xlsx_response.content), data_only=False)
    sheet = workbook["内容"]
    values = {
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert {"事项", "负责人", "整理进展", "小林"}.issubset(values)
    assert sheet.freeze_panes == "A3"

    other = client.post(
        "/api/v1/auth/register",
        json={
            "email": "other@example.com",
            "password": "safe-password-456",
            "display_name": "其他同学",
        },
    ).json()
    forbidden = client.get(
        f"/api/v1/artifacts/{by_type['docx']['id']}/download",
        headers={"Authorization": f"Bearer {other['access_token']}"},
    )
    assert forbidden.status_code == 404
    assert forbidden.json()["error"]["code"] == "ARTIFACT_NOT_FOUND"


def test_chat_plain_word_request_tells_model_to_create_and_returns_download(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    async def fake_stream(self, messages):
        del self
        system_prompt = messages[0]["content"]
        assert "具备生成并提供可下载的 Word（DOCX）文件的能力" in system_prompt
        assert "不得声称无法生成" in system_prompt
        yield "# 通用工作文档\n\n## 目标\n\n请在此填写目标。\n\n## 内容\n\n请在此补充正文。"

    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "Word 创作", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "你给我生成份word文档"},
    )

    assert response.status_code == 200
    assert "event: artifacts" in response.text
    assert "你希望把本次内容整理成Word 文档并直接下载" in response.text
    assert "我会整理适合写入文件的完整内容，并生成可下载的Word 文档" in response.text
    assert "当前环境无法生成或发送文件" not in response.text
    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    artifact = messages[-1]["generated_artifacts"][0]
    assert artifact["artifact_type"] == "docx"
    downloaded = client.get(
        f"/api/v1/artifacts/{artifact['id']}/download", headers=auth_headers
    )
    document = Document(BytesIO(downloaded.content))
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "通用工作文档" in document_text
    assert "无法生成" not in document_text


def test_chat_can_reformat_latest_generated_docx_and_return_new_download(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from docx.oxml.ns import qn

    from app.infrastructure.llm.deepseek import DeepSeekProvider

    async def fake_stream(self, messages):
        del self, messages
        yield "# 面经报告\n\n## 技术总结\n\n这里是报告正文。"

    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "文档格式修改", "mode": "standard"},
    ).json()
    first = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "请生成Word文档：面经报告"},
    )
    assert first.status_code == 200
    assert "event: artifacts" in first.text

    reformatted = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "把字体颜色统一一下，给我返回生成后的文档"},
    )
    assert reformatted.status_code == 200
    assert "event: artifacts" in reformatted.text

    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    artifact = messages[-1]["generated_artifacts"][0]
    assert artifact["filename"].endswith("_黑色.docx")
    download = client.get(
        f"/api/v1/artifacts/{artifact['id']}/download", headers=auth_headers
    )
    document = Document(BytesIO(download.content))
    text_runs = [
        run
        for run in document.element.body.iter(qn("w:r"))
        if any(node.text for node in run.iter(qn("w:t")))
    ]
    assert text_runs
    assert all(
        run.get_or_add_rPr().find(qn("w:color")).get(qn("w:val")) == "000000"
        for run in text_runs
    )


def test_chat_edits_uploaded_docx_and_returns_preserved_download(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    template = Document()
    table = template.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "实习主要工作内容"
    answer_run = table.cell(0, 1).paragraphs[0].add_run("待填写")
    answer_run.bold = True
    table.cell(1, 0).text = "联系邮箱"
    table.cell(1, 1).text = "student@example.com"
    payload = BytesIO()
    template.save(payload)
    material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "实习问卷.docx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"purpose": "conversation"},
    ).json()

    async def fake_complete_json(self, messages, *, max_tokens=2400):
        del self, max_tokens
        serialized_messages = str(messages)
        assert "student@example.com" not in serialized_messages
        assert "[邮箱已脱敏]" in serialized_messages
        return {
            "message": "已根据现有内容补全主观题。",
            "missing_information": [],
            "changes": [
                {
                    "location": "table:0:row:0:cell:1",
                    "expected_text": "待填写",
                    "value": "协助整理需求、分析数据并完成周报。",
                }
            ],
        }

    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_complete_json)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "补全问卷", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "这个补全发给我", "attachment_ids": [material["id"]]},
    )
    assert response.status_code == 200
    assert "event: artifacts" in response.text
    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    artifact = messages[-1]["generated_artifacts"][0]
    assert artifact["filename"] == "实习问卷_已补全.docx"
    download = client.get(
        f"/api/v1/artifacts/{artifact['id']}/download", headers=auth_headers
    )
    edited = Document(BytesIO(download.content))
    answer = edited.tables[0].cell(0, 1)
    assert answer.text == "协助整理需求、分析数据并完成周报。"
    assert answer.paragraphs[0].runs[0].bold is True


def test_chat_edits_uploaded_xlsx_without_overwriting_format_or_formula(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    template = Workbook()
    sheet = template.active
    sheet.title = "登记表"
    sheet["A1"] = "本周进展"
    sheet["B1"] = "待填写"
    sheet["B1"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["C1"] = "=1+1"
    payload = BytesIO()
    template.save(payload)
    material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={
            "file": (
                "周报.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"purpose": "conversation"},
    ).json()

    async def fake_complete_json(self, messages, *, max_tokens=2400):
        del self, messages, max_tokens
        return {
            "message": "已更新表格。",
            "missing_information": [],
            "changes": [
                {
                    "sheet": "登记表",
                    "cell": "B1",
                    "expected_text": "待填写",
                    "value": "完成需求评审和竞品分析",
                },
                {
                    "sheet": "登记表",
                    "cell": "C1",
                    "expected_text": "=1+1",
                    "value": "禁止覆盖公式",
                },
            ],
        }

    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_complete_json)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "更新表格", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "请补全这个表格并发给我", "attachment_ids": [material["id"]]},
    )
    assert response.status_code == 200
    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    artifact = messages[-1]["generated_artifacts"][0]
    download = client.get(
        f"/api/v1/artifacts/{artifact['id']}/download", headers=auth_headers
    )
    edited = load_workbook(BytesIO(download.content), data_only=False)
    edited_sheet = edited["登记表"]
    assert edited_sheet["B1"].value == "完成需求评审和竞品分析"
    assert edited_sheet["B1"].fill.fgColor.rgb == "00FFF2CC"
    assert edited_sheet["C1"].value == "=1+1"


def test_chat_asks_for_missing_facts_before_editing_attachment(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    template = Document()
    table = template.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "实习单位名称"
    table.cell(0, 1).text = "待填写"
    payload = BytesIO()
    template.save(payload)
    material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("实习信息.docx", payload.getvalue(), "application/octet-stream")},
        data={"purpose": "conversation"},
    ).json()

    async def fake_complete_json(self, messages, *, max_tokens=2400):
        del self, max_tokens
        if "用户补充：实习单位名称是启程科技" in str(messages):
            return {
                "message": "已根据补充信息完成填写。",
                "missing_information": [],
                "changes": [
                    {
                        "location": "table:0:row:0:cell:1",
                        "expected_text": "待填写",
                        "value": "启程科技",
                    }
                ],
            }
        return {
            "message": "补全前需要确认客观信息。",
            "missing_information": ["实习单位名称"],
            "changes": [],
        }

    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_complete_json)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "补全实习信息", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "补全后发给我", "attachment_ids": [material["id"]]},
    )
    assert response.status_code == 200
    assert "实习单位名称" in response.text
    assert "event: artifacts" not in response.text
    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    assert messages[-1]["generated_artifacts"] == []

    resumed = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "实习单位名称是启程科技"},
    )
    assert resumed.status_code == 200
    assert "event: artifacts" in resumed.text
    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    artifact = messages[-1]["generated_artifacts"][0]
    download = client.get(
        f"/api/v1/artifacts/{artifact['id']}/download", headers=auth_headers
    )
    edited = Document(BytesIO(download.content))
    assert edited.tables[0].cell(0, 1).text == "启程科技"


def test_chat_keeps_unknown_fields_blank_and_still_returns_attachment(
    client: TestClient, auth_headers: dict[str, str], monkeypatch
) -> None:
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    template = Document()
    table = template.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "单位负责人签名"
    table.cell(0, 1).text = ""
    payload = BytesIO()
    template.save(payload)
    material = client.post(
        "/api/v1/materials",
        headers=auth_headers,
        files={"file": ("实习问卷.docx", payload.getvalue(), "application/octet-stream")},
        data={"purpose": "conversation"},
    ).json()

    async def fake_complete_json(self, messages, *, max_tokens=2400):
        del self, messages, max_tokens
        return {
            "message": "仍缺少单位负责人签名。",
            "missing_information": ["单位负责人签名"],
            "changes": [],
        }

    monkeypatch.setattr(DeepSeekProvider, "complete_json", fake_complete_json)
    conversation = client.post(
        "/api/v1/conversations",
        headers=auth_headers,
        json={"title": "保留空白", "mode": "standard"},
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={
            "content": "帮我补全，不清楚的地方可以先不填写，补全后发给我",
            "attachment_ids": [material["id"]],
        },
    )
    assert response.status_code == 200
    assert "event: artifacts" in response.text
    assert "保留无法确认的字段" in response.text
    messages = client.get(
        f"/api/v1/conversations/{conversation['id']}/messages", headers=auth_headers
    ).json()
    artifact = messages[-1]["generated_artifacts"][0]
    assert artifact["filename"] == "实习问卷_已补全.docx"
