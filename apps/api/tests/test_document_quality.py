import asyncio
from io import BytesIO

import httpx
import pytest
from docx import Document
from docx.oxml.ns import qn
from pydantic import ValidationError

from app.core.config import Settings
from app.core.errors import AppError
from app.modules.artifacts.quality import inspect_docx, render_check, review_document
from app.modules.artifacts.service import build_docx
from app.tools.factory import build_tool_registry
from app.tools.web_fetch import WebFetchHandler, WebFetchInput

DRAFT = """# 活动安排

## 上午活动

**地点：** 市图书馆，阅读[主办方公告](https://example.org/event)。

| 活动 | 时间 |
| --- | --- |
| 咨询 | 09:30 |

## 下午活动

> 场次以官方公告为准。

1. 查票
2. 到场

| 项目 | 说明 |
| --- | --- |
| 展览 | 免费 |

## 来源说明

1. 原始公告
"""


def test_docx_preserves_order_multiple_tables_and_inline_styles():
    _, binary = build_docx("生成文档", DRAFT)
    document = Document(BytesIO(binary))
    blocks = list(document.element.body)
    sequence = [(block.tag.rsplit("}", 1)[-1], "".join(block.itertext())) for block in blocks]
    first_table = next(i for i, (kind, _) in enumerate(sequence) if kind == "tbl")
    afternoon = next(i for i, (_, text) in enumerate(sequence) if "下午活动" in text)
    sources = next(i for i, (_, text) in enumerate(sequence) if "来源说明" in text)
    assert first_table < afternoon < sources
    assert len(document.tables) == 2
    assert document.tables[1].cell(1, 0).text == "展览"
    paragraph = next(p for p in document.paragraphs if "地点" in p.text)
    assert any(r.bold and r.text == "地点：" for r in paragraph.runs)
    assert "**" not in paragraph.text
    assert paragraph.hyperlinks[0].url == "https://example.org/event"
    assert not any(p.text.startswith(">") for p in document.paragraphs)
    numbered = [p for p in document.paragraphs if p.text in {"查票", "原始公告"}]
    assert len({p._p.pPr.numPr.numId.val for p in numbered}) == 2
    assert document.paragraphs[0].style.name == "Title"
    assert document.styles["Title"]._element.pPr.find(qn("w:pBdr")) is None
    numbering_tags = [element.tag for element in document.part.numbering_part.element]
    first_number = numbering_tags.index(qn("w:num"))
    assert qn("w:abstractNum") not in numbering_tags[first_number:]
    assert document.sections[0].page_width > 0
    for table in document.tables:
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert inspect_docx(binary, DRAFT) == []


def test_docx_literal_pipes_and_escaped_markdown_are_not_dropped():
    _, binary = build_docx(
        "生成文档",
        "# 标题\n\n正文 A | B\n\n\\*原样星号\\*\n\n"
        "| 列一 | 列二 |\n| --- | --- |\n| A \\| B | **重点** |",
    )
    doc = Document(BytesIO(binary))
    assert "正文 A | B" in [p.text for p in doc.paragraphs]
    assert "*原样星号*" in [p.text for p in doc.paragraphs]
    assert doc.tables[0].cell(1, 0).text == "A | B"
    assert any(run.bold for run in doc.tables[0].cell(1, 1).paragraphs[0].runs)


class ReviewProvider:
    def __init__(self, results):
        self.results = iter(results)
        self.calls = 0

    async def complete_json(self, messages, *, max_tokens):
        assert "untrusted_sources" in messages[-1]["content"]
        self.calls += 1
        result = next(self.results)
        if isinstance(result, Exception):
            raise result
        return result


def run_review(provider, settings):
    async def run():
        return [
            event
            async for event in review_document(
                provider=provider,
                settings=settings,
                prompt="生成活动 Word 文档",
                content=DRAFT,
                sources=[],
            )
        ]

    return asyncio.run(run())


def test_document_review_rebuilds_and_rechecks_revised_document():
    revised = DRAFT.replace("免费", "是否收费待核实")
    provider = ReviewProvider(
        [
            {"approved": False, "issues": ["费用缺少来源"], "revised_content": revised},
            {"approved": True},
        ]
    )
    events = run_review(
        provider,
        Settings(
            _env_file=None, demo_mode=False, deepseek_api_key="test", document_render_enabled=False
        ),
    )
    final = events[-1]
    assert final["content"] == revised.strip()
    assert final["report"]["revisions"] == 1
    assert final["report"]["model_review"] == "passed"
    assert provider.calls == 2
    assert sum(e["stage"] == "checking" for e in events) == 2


@pytest.mark.parametrize(
    "result",
    [
        AppError("MODEL_PROVIDER_ERROR", "unavailable", status_code=502),
        {"understanding": "malformed review"},
        {"approved": False, "issues": ["too brief"], "revised_content": "# 不完整"},
    ],
)
def test_review_failure_preserves_complete_draft(result):
    final = run_review(
        ReviewProvider([result]),
        Settings(
            _env_file=None, demo_mode=False, deepseek_api_key="test", document_render_enabled=False
        ),
    )[-1]
    assert final["content"] == DRAFT
    assert final["report"]["warnings"]


def test_review_revision_limit_and_missing_renderer_are_honest():
    provider = ReviewProvider([{"approved": False, "issues": ["未核验"], "revised_content": DRAFT}])
    final = run_review(
        provider,
        Settings(
            _env_file=None,
            demo_mode=False,
            deepseek_api_key="test",
            document_review_max_revisions=0,
            document_render_enabled=False,
        ),
    )[-1]
    assert provider.calls == 1
    assert final["report"]["model_review"] == "warning"
    assert final["report"]["render"]["status"] == "skipped"


def test_local_tools_survive_without_web_search_credentials():
    registry = build_tool_registry(None, Settings(_env_file=None, qdrant_enabled=False))
    schemas = registry.model_schemas(
        effects={"read", "draft"}, excluded_names={"web.search", "web.fetch"}
    )
    names = {item["function"]["name"] for item in schemas}
    assert {"resume__search", "recruitment__search"}.issubset(names)
    assert "web__search" not in names


@pytest.mark.parametrize(
    "url",
    [
        "http://localhost/",
        "http://127.0.0.1/",
        "http://[::1]/",
        "http://192.168.1.1/",
        "https://user:pass@example.com/",
    ],
)
def test_web_fetch_rejects_local_or_credential_urls(url):
    with pytest.raises(ValidationError):
        WebFetchInput(url=url)


def test_web_fetch_uses_provider_extract_and_bounds_content(monkeypatch):
    requests = []

    def handler(request):
        requests.append(request)
        return httpx.Response(
            200,
            json={
                "results": [
                    {"url": "https://example.org/", "title": "公告", "raw_content": "正文" * 1000}
                ]
            },
        )

    factory = httpx.AsyncClient
    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: factory(transport=httpx.MockTransport(handler), **kwargs),
    )
    result = asyncio.run(
        WebFetchHandler(Settings(_env_file=None, tavily_api_key="test"))(
            WebFetchInput(url="https://example.org/", max_chars=1000), None
        )
    )
    assert result.truncated
    assert len(result.content) == 1000
    assert requests[0].url.path == "/extract"
    assert requests[0].url.host == "api.tavily.com"


def test_missing_renderer_never_reports_passed(monkeypatch):
    monkeypatch.setattr("app.modules.artifacts.quality.shutil.which", lambda name: None)
    result = render_check(b"unused", Settings(_env_file=None))
    assert result["status"] == "unavailable"


def test_pdf_check_reports_blank_pages(monkeypatch):
    from pathlib import Path

    import pymupdf

    def fake_convert(command, **kwargs):
        folder = Path(command[command.index("--outdir") + 1])
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.save(folder / "document.pdf")

    monkeypatch.setattr("app.modules.artifacts.quality.subprocess.run", fake_convert)
    result = render_check(b"test", Settings(_env_file=None, document_renderer_path="test-renderer"))
    assert result["status"] == "warning"
    assert result["pages"] == 1
    assert result["issues"] == ["第 1 页为空白页"]


def test_local_tools_are_offered_in_chat_when_web_is_unconfigured(
    client,
    auth_headers,
    monkeypatch,
):
    from app.infrastructure.llm.deepseek import DeepSeekProvider

    offered = set()

    async def select(self, messages, tools, **kwargs):
        offered.update(tool["function"]["name"] for tool in tools)
        return None, []

    async def narrative(self, messages, **kwargs):
        return {"understanding": "检查简历经历", "next_action": "我会读取简历。"}

    async def stream(self, messages):
        yield "可继续查阅简历和本地职位。"

    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(Settings, "online_search_available", property(lambda self: False))
    monkeypatch.setattr(DeepSeekProvider, "select_tool_calls", select)
    monkeypatch.setattr(DeepSeekProvider, "complete_json", narrative)
    monkeypatch.setattr(DeepSeekProvider, "stream_chat", stream)
    conversation = client.post(
        "/api/v1/conversations", headers=auth_headers, json={"title": "本地检索"}
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "帮我分析简历经历"},
    )
    assert "event: message.completed" in response.text
    assert {"resume__search", "recruitment__search"}.issubset(offered)
    assert "web__search" not in offered
    assert "web__fetch" not in offered
