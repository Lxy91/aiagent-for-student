from io import BytesIO

import pytest
from docx import Document

from app.core.config import Settings, get_settings
from app.infrastructure.llm.deepseek import DeepSeekProvider


@pytest.fixture(autouse=True)
def semantic_provider(monkeypatch):
    # Mock only the remote model boundary; validation/persistence stay real.
    async def complete(self, messages, **kwargs):
        import json

        source = json.loads(messages[1]["content"])["resume_document"]
        result = {}
        for key, value in [
            ("identity", "React 项目负责人，三年经验"),
            ("identity", "王同学"),
            ("identity", "前端开发"),
            ("education", "示例大学 本科"),
            ("skills", "React TypeScript"),
            ("projects", "星河平台"),
        ]:
            if value in source:
                result.setdefault(key, []).append({"text": value, "evidence": [value]})
        return result

    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(DeepSeekProvider, "complete_json", complete)


def upload(client, headers, text="React 项目负责人，三年经验"):
    return client.put(
        "/api/v1/resume",
        headers=headers,
        files={"file": ("resume.txt", text.encode(), "text/plain")},
    )


def test_resume_lifecycle_and_isolation(client, auth_headers):
    assert client.get("/api/v1/resume", headers=auth_headers).json() is None
    result = upload(client, auth_headers)
    assert result.status_code == 200
    assert result.json()["active"] is False
    assert "React" in result.json()["content"]
    assert client.patch("/api/v1/resume", headers=auth_headers, json={"active": True}).json()[
        "active"
    ]
    other = client.post(
        "/api/v1/auth/register",
        json={
            "email": "other@example.com",
            "password": "safe-password-123",
            "display_name": "Other",
        },
    ).json()
    headers = {"Authorization": f"Bearer {other['access_token']}"}
    assert client.get("/api/v1/resume", headers=headers).json() is None
    assert client.patch("/api/v1/resume", headers=headers, json={"active": True}).status_code == 404
    assert client.delete("/api/v1/resume", headers=headers).status_code == 204
    assert client.get("/api/v1/resume", headers=auth_headers).json()["active"]
    assert upload(client, auth_headers, "新简历").json()["active"] is False
    assert client.delete("/api/v1/resume", headers=auth_headers).status_code == 204
    assert client.get("/api/v1/resume", headers=auth_headers).json() is None


def test_resume_validation_and_docx(client, auth_headers):
    for filename, data, expected in [
        ("bad.exe", b"test", 415),
        ("empty.txt", b"", 400),
        ("bad.docx", b"bad", 422),
        ("long.txt", b"a" * 30001, 422),
    ]:
        response = client.put(
            "/api/v1/resume", headers=auth_headers, files={"file": (filename, data)}
        )
        assert response.status_code == expected
    doc = Document()
    doc.add_paragraph("Resume document experience")
    buffer = BytesIO()
    doc.save(buffer)
    result = client.put(
        "/api/v1/resume", headers=auth_headers, files={"file": ("resume.docx", buffer.getvalue())}
    )
    assert result.status_code == 200
    assert "Resume document experience" in result.json()["content"]


def test_resume_chat_opt_in(client, auth_headers, monkeypatch):
    captured = []

    async def fake_stream(self, messages, tools=None):
        captured.append(messages)
        yield "回答"

    upload(client, auth_headers)
    monkeypatch.setattr(DeepSeekProvider, "stream_chat", fake_stream)
    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: False))
    monkeypatch.setattr(Settings, "online_search_available", property(lambda self: False))
    for active in [False, True, False]:
        client.patch("/api/v1/resume", headers=auth_headers, json={"active": active})
        conversation = client.post(
            "/api/v1/conversations", headers=auth_headers, json={"mode": "standard"}
        ).json()
        response = client.post(
            f"/api/v1/conversations/{conversation['id']}/messages:stream",
            headers=auth_headers,
            json={"content": "你好"},
        )
        assert response.status_code == 200
        assert ("React 项目负责人" in captured[-1][0]["content"]) is active


def test_resume_interview_opt_in(client, auth_headers, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "volc_rtc_app_id", "123456789012345678901234")
    monkeypatch.setattr(settings, "volc_rtc_app_key", "test-secret")
    monkeypatch.setattr(settings, "volc_access_key_id", "test-key")
    monkeypatch.setattr(settings, "volc_secret_access_key", "test-secret")
    upload(client, auth_headers)
    for active in [False, True, False]:
        client.patch("/api/v1/resume", headers=auth_headers, json={"active": active})
        response = client.post(
            "/api/v1/interviews/sessions",
            headers=auth_headers,
            json={"track": "react", "enable_avatar": False},
        )
        assert response.status_code == 201
        assert ("React 项目负责人" in response.json()["interview_prompt"]) is active


def test_summary_and_detail_search_obey_consent(client, auth_headers):
    text = (
        "王同学，前端工程师\n教育背景\n示例大学 本科\n技能\nReact TypeScript\n项目经历\n星河平台\n"
        + "星河使用自研调度器降低渲染延迟，" * 60
    )
    result = upload(client, auth_headers, text)
    assert result.status_code == 200
    summary = result.json()["summary"]
    assert len(summary) <= 900
    assert "示例大学" in summary and "React" in summary
    assert "自研调度器" not in summary
    request = {"arguments": {"query": "星河调度器"}}
    url = "/api/v1/tools/resume.search:invoke"
    assert client.post(url, headers=auth_headers, json=request).json()["result"]["results"] == []
    client.patch("/api/v1/resume", headers=auth_headers, json={"active": True})
    matches = client.post(url, headers=auth_headers, json=request).json()["result"]["results"]
    assert matches and any("自研调度器" in part for part in matches)
    assert sum(map(len, matches)) <= 2400
    no_match = client.post(url, headers=auth_headers, json={"arguments": {"query": "unrelatedXYZ"}})
    assert no_match.json()["result"]["results"] == []
    client.patch("/api/v1/resume", headers=auth_headers, json={"active": False})
    assert client.post(url, headers=auth_headers, json=request).json()["result"]["results"] == []


def test_model_looks_up_resume_only_when_needed(client, auth_headers, monkeypatch):
    text = "前端开发\n技能\nReact\n项目经历\n星河\n" + "星河实现了可配置调度器，" * 60
    upload(client, auth_headers, text)
    client.patch("/api/v1/resume", headers=auth_headers, json={"active": True})
    captured = []

    async def select(self, messages, tools, **kwargs):
        if messages[-1]["role"] == "tool":
            return None, []
        return {
            "role": "assistant",
            "content": None,
            "tool_calls": [
                {
                    "id": "resume-detail",
                    "type": "function",
                    "function": {"name": "resume__search", "arguments": '{"query":"星河调度器"}'},
                }
            ],
        }, [
            {
                "id": "resume-detail",
                "type": "function",
                "function": {"name": "resume__search", "arguments": '{"query":"星河调度器"}'},
            }
        ]

    async def narrative(self, messages, **kwargs):
        return {"understanding": "查阅项目细节", "next_action": "查阅简历相关片段"}

    async def stream(self, messages):
        captured.extend(messages)
        yield "项目中使用了可配置调度器。"

    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: True))
    monkeypatch.setattr(Settings, "online_search_available", property(lambda self: False))
    monkeypatch.setattr(DeepSeekProvider, "select_tool_calls", select)
    monkeypatch.setattr(DeepSeekProvider, "complete_json", narrative)
    monkeypatch.setattr(DeepSeekProvider, "stream_chat", stream)
    conversation = client.post(
        "/api/v1/conversations", headers=auth_headers, json={"mode": "standard"}
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages:stream",
        headers=auth_headers,
        json={"content": "详细介绍一下我简历中的星河调度器"},
    )
    assert response.status_code == 200
    assert '"tool": "resume.search"' in response.text
    assert "可配置调度器" not in captured[0]["content"]
    assert any(
        "可配置调度器" in str(item["content"]) for item in captured if item["role"] == "tool"
    )


def test_failed_summary_preserves_source_without_fake_fallback(client, auth_headers, monkeypatch):
    async def unsupported(self, messages, **kwargs):
        return {"education": [{"text": "不存在的大学", "evidence": ["伪造的原文"]}]}

    monkeypatch.setattr(DeepSeekProvider, "complete_json", unsupported)
    result = upload(client, auth_headers)
    assert result.status_code == 200
    assert result.json()["content"] == "React 项目负责人，三年经验"
    assert result.json()["summary"] == ""
    assert result.json()["summary_error"]
    monkeypatch.setattr(Settings, "deepseek_enabled", property(lambda self: False))
    regenerated = client.post("/api/v1/resume:regenerate-summary", headers=auth_headers)
    assert regenerated.status_code == 200
    assert regenerated.json()["summary"] == ""


def test_legacy_summary_rebuilt_once_without_changing_consent(client, auth_headers, monkeypatch):
    import asyncio

    from conftest import test_session_factory
    from sqlalchemy import select

    from app.infrastructure.db.models import ResumeModel

    upload(client, auth_headers)
    client.patch("/api/v1/resume", headers=auth_headers, json={"active": True})

    async def make_legacy():
        async with test_session_factory() as session:
            item = await session.scalar(select(ResumeModel))
            item.summary = "基本信息：Personal resume；细心从每一个小细节开始"
            item.summary_version = 0
            await session.commit()

    asyncio.run(make_legacy())
    calls = []

    async def complete(self, messages, **kwargs):
        calls.append(messages)
        return {"skills": [{"text": "React", "evidence": ["React"]}]}

    monkeypatch.setattr(DeepSeekProvider, "complete_json", complete)
    result = client.get("/api/v1/resume", headers=auth_headers).json()
    assert result["active"] is True
    assert result["summary"] == "技能：React"
    assert "Personal resume" not in result["summary"]
    client.get("/api/v1/resume", headers=auth_headers)
    assert len(calls) == 1


def test_original_attachment_download_and_user_isolation(client, auth_headers):
    content = "React 项目负责人，三年经验".encode()
    assert client.get("/api/v1/resume/attachment").status_code == 401
    assert client.get("/api/v1/resume/attachment", headers=auth_headers).status_code == 404
    assert upload(client, auth_headers).json()["has_attachment"] is True
    response = client.get("/api/v1/resume/attachment", headers=auth_headers)
    assert response.status_code == 200
    assert response.content == content
    assert response.headers["content-type"].startswith("text/plain")
    assert response.headers["content-disposition"].startswith("inline;")
    assert response.headers["cache-control"] == "private, no-store"
    user = client.post(
        "/api/v1/auth/register",
        json={
            "email": "attachment@example.com",
            "password": "safe-password-123",
            "display_name": "Other",
        },
    ).json()
    other = {"Authorization": f"Bearer {user['access_token']}"}
    assert client.get("/api/v1/resume/attachment", headers=other).status_code == 404
    upload(client, auth_headers, "替换后的简历")
    assert (
        client.get("/api/v1/resume/attachment", headers=auth_headers).content
        == "替换后的简历".encode()
    )
    client.delete("/api/v1/resume", headers=auth_headers)
    assert client.get("/api/v1/resume/attachment", headers=auth_headers).status_code == 404


def test_pdf_docx_attachments_preserve_original_bytes(client, auth_headers):
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((60, 60), "Original resume PDF")
    pdf_bytes = pdf.tobytes()
    pdf.close()
    doc = Document()
    doc.add_heading("Original resume Word", level=1)
    buffer = BytesIO()
    doc.save(buffer)
    for filename, data, mime in [
        ("中文简历.pdf", pdf_bytes, "application/pdf"),
        (
            "中文简历.docx",
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
    ]:
        response = client.put(
            "/api/v1/resume", headers=auth_headers, files={"file": (filename, data)}
        )
        assert response.status_code == 200
        assert response.json()["has_attachment"] is True
        attachment = client.get("/api/v1/resume/attachment", headers=auth_headers)
        assert attachment.content == data
        assert attachment.headers["content-type"] == mime
        assert "%E4%B8%AD%E6%96%87" in attachment.headers["content-disposition"]


def test_legacy_resume_reports_missing_original_attachment(client, auth_headers):
    import asyncio

    from conftest import test_session_factory
    from sqlalchemy import select

    from app.infrastructure.db.models import ResumeModel

    upload(client, auth_headers)

    async def legacy():
        async with test_session_factory() as session:
            item = await session.scalar(select(ResumeModel))
            item.attachment_size = None
            item.binary_content = None
            await session.commit()

    asyncio.run(legacy())
    result = client.get("/api/v1/resume", headers=auth_headers)
    assert result.json()["has_attachment"] is False
    assert result.json()["content"]
    response = client.get("/api/v1/resume/attachment", headers=auth_headers)
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "RESUME_ATTACHMENT_MISSING"
